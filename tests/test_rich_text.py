"""Shared rich-value behavior and reference reopening of canonical output."""

import importlib
import zipfile

import crabxl
import openpyxl
import pytest


@pytest.mark.parametrize("engine", [openpyxl, crabxl])
def test_rich_runs_live_mutation_loading_and_repeat_output(
    engine, tmp_path, shared_strings_source
):
    rich_module = importlib.import_module(engine.__name__ + ".cell.rich_text")
    InlineFont = importlib.import_module(engine.__name__ + ".cell.text").InlineFont
    CellRichText, TextBlock = rich_module.CellRichText, rich_module.TextBlock
    for loaded in (False, True):
        book = engine.Workbook()
        sheet = book.active
        value = CellRichText(
            " =literal ",
            TextBlock(InlineFont(rFont="Arial", b=True, color="abCDef"), "styled"),
            " tail ",
        )
        sheet["A1"] = value
        assert sheet["A1"].data_type == "s"
        if loaded:
            source = tmp_path / f"{engine.__name__}-rich-source.xlsx"
            book.save(source)
            book.close()
            plain = engine.load_workbook(source)
            assert plain.active["A1"].value == str(value)
            plain.close()
            book = engine.load_workbook(source, rich_text=True)
            sheet = book.active
            if engine is crabxl:
                untouched = tmp_path / "rich-untouched.xlsx"
                book.save(untouched)
                with zipfile.ZipFile(source) as a, zipfile.ZipFile(untouched) as b:
                    assert {n: a.read(n) for n in a.namelist()} == {
                        n: b.read(n) for n in b.namelist()
                    }
            value = sheet["A1"].value
        assert isinstance(value, CellRichText)
        assert value[1].font.rFont == "Arial"
        assert value[1].font.color.rgb == "00abCDef"
        assert sheet["A1"].value is value
        sheet["B1"] = value
        value[1].text = "updated"
        value[1].font.i = True
        value[1].font.color.tint = 0.25
        value.append("end")
        assert (
            str(sheet["A1"].value)
            == str(sheet["B1"].value)
            == " =literal updated tail end"
        )
        assert next(sheet.iter_rows(max_row=1, max_col=1, values_only=True))[0] is value
        sheet.insert_rows(1)
        value[1].text = "moved"
        assert (
            str(sheet["A2"].value)
            == str(sheet["B2"].value)
            == " =literal moved tail end"
        )
        for suffix in (1, 2):
            output = tmp_path / f"{engine.__name__}-rich-{loaded}-{suffix}.xlsx"
            book.save(output)
            reference = openpyxl.load_workbook(output, rich_text=True)
            rich = reference.active["A2"].value
            assert str(rich) == " =literal moved tail end"
            assert rich[1].font.b and rich[1].font.i
            assert rich[1].font.color.tint == 0.25
            reference.close()
            for values_only in (False, True):
                reader = engine.load_workbook(output, read_only=True, rich_text=True)
                row = next(
                    reader.active.iter_rows(
                        min_row=2, max_row=2, max_col=1, values_only=values_only
                    )
                )
                rich = row[0] if values_only else row[0].value
                assert isinstance(rich, str) and rich == " =literal moved tail end"
                if not values_only:
                    assert row[0].data_type == "s"
                reader.close()
        sheet["A2"] = "ordinary"
        value[1].text = "remaining"
        assert sheet["A2"].value == "ordinary"
        assert str(sheet["B2"].value) == " =literal remaining tail end"
        book.close()

    # Shared rich strings retain typed read-only values in the reference API.
    source = shared_strings_source(
        tmp_path / f"{engine.__name__}-shared-rich.xlsx", ["placeholder"]
    )
    with zipfile.ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/sharedStrings.xml"] = (
        b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="1" uniqueCount="1">'
        b'<si><r><rPr><rFont val="Arial"/><b/><color rgb="00abCDef"/></rPr><t xml:space="preserve"> rich </t></r>'
        b'<r><t>tail</t></r><rPh sb="0" eb="1"><t>pronunciation</t></rPh>'
        b'<phoneticPr fontId="0" type="Hiragana" alignment="center"/></si></sst>'
    )
    with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    for readonly in (False, True):
        options = {}
        if engine is crabxl:
            options["resource_options"] = crabxl.ResourceOptions(
                shared_strings=crabxl.SharedStringOptions(storage="disk")
            )
        book = engine.load_workbook(
            source, rich_text=True, read_only=readonly, **options
        )
        value = next(book.active.iter_rows(max_col=1, values_only=True))[0]
        assert isinstance(value, CellRichText) and str(value) == " rich tail"
        assert value[0].font.b and value[0].font.rFont == "Arial"
        if not readonly:
            value[0].text = "changed"
            book.active.insert_rows(1)
            output = tmp_path / f"{engine.__name__}-shared-rich-edit.xlsx"
            book.save(output)
            reference = openpyxl.load_workbook(output, rich_text=True)
            assert str(reference.active["A2"].value) == "changedtail"
            reference.close()
            if engine is crabxl:
                with zipfile.ZipFile(output) as archive:
                    xml = archive.read("xl/worksheets/sheet1.xml")
                    assert (
                        b'<rPh sb="0" eb="1">' in xml
                        and b'<phoneticPr fontId="0"' in xml
                    )
        book.close()

    stream_book = engine.Workbook(write_only=True)
    stream_sheet = stream_book.create_sheet("Stream")
    stream_value = CellRichText("prefix", TextBlock(InlineFont(b=True), "bold"))
    stream_sheet.append([stream_value])
    stream_value[1].text = "after flush"
    output = tmp_path / f"{engine.__name__}-stream-rich.xlsx"
    stream_book.save(output)
    reference = openpyxl.load_workbook(output, rich_text=True)
    assert str(reference.active["A1"].value) == "prefixbold"
    assert reference.active["A1"].value[1].font.b
    reference.close()
    stream_book.close()


def test_rich_mutation_budget_failure_keeps_all_owners_and_view_atomic():
    from crabxl.cell.rich_text import CellRichText, TextBlock
    from crabxl.cell.text import InlineFont

    large = crabxl.Workbook(max_memory_bytes=256 * 1024)
    small = crabxl.Workbook(max_memory_bytes=32 * 1024)
    value = CellRichText(TextBlock(InlineFont(b=True), "original"))
    large.active["A1"] = value
    small.active["A1"] = value
    with pytest.raises(MemoryError):
        value[0].text = "x" * (64 * 1024)
    assert (
        str(value)
        == str(large.active["A1"].value)
        == str(small.active["A1"].value)
        == "original"
    )
    value[0].text = "retry"
    assert str(large.active["A1"].value) == str(small.active["A1"].value) == "retry"
    large.close()
    small.close()
