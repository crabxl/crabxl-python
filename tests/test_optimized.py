"""Public optimized-mode parity and deterministic resource ownership checks."""

import os
import zipfile
from datetime import date, datetime, time, timedelta
from pathlib import Path

import crabxl
import openpyxl
import pytest


@pytest.fixture(params=[openpyxl, crabxl], ids=["openpyxl", "crabxl"])
def engine(request):
    return request.param


def source(path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(
        [
            1,
            None,
            "=A1+1",
            date(2024, 2, 29),
            time(12, 3, 4),
            timedelta(days=2, seconds=3),
            True,
            " text ",
            "#N/A",
        ]
    )
    sheet["B3"] = "middle"
    sheet["D5"] = "last"
    book.create_sheet("Other").append([42])
    book.save(path)
    book.close()


@pytest.mark.parametrize("cached", [False, True])
def test_read_only_values_cells_sparse_bounds_and_dates(engine, cached, tmp_path):
    path = tmp_path / "read.xlsx"
    source(path)
    book = engine.load_workbook(path, read_only=True, data_only=cached)
    assert book.read_only and not book.write_only
    assert book.sheetnames == ["Sheet", "Other"]
    sheet = book.active
    assert (sheet.min_row, sheet.min_column, sheet.max_row, sheet.max_column) == (
        1,
        1,
        5,
        9,
    )
    assert sheet.calculate_dimension() == "A1:I5"
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == (
        1,
        None,
        None if cached else "=A1+1",
        datetime(2024, 2, 29),
        time(12, 3, 4),
        timedelta(days=2, seconds=3),
        True,
        " text ",
        "#N/A",
    )
    assert rows[1] == (None,) * 9
    assert rows[2][1] == "middle" and rows[4][3] == "last"
    assert list(
        sheet.iter_rows(min_row=2, max_row=3, min_col=2, max_col=3, values_only=True)
    ) == [(None, None), ("middle", None)]
    cell = sheet["D1"]
    assert cell.coordinate == "D1" and cell.column_letter == "D"
    assert cell.is_date and cell.has_style and cell.number_format == "yyyy-mm-dd"
    with pytest.raises(AttributeError):
        cell.value = 2
    with pytest.raises(TypeError):
        book.save(tmp_path / "invalid.xlsx")
    book.close()


@pytest.mark.parametrize("dimension", [None, "A1:A1"])
def test_read_only_dimensions_can_be_reset_and_calculated(engine, dimension, tmp_path):
    path = tmp_path / "dimensions.xlsx"
    source(path)
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    old = b'<dimension ref="A1:I5"/>'
    replacement = (
        b"" if dimension is None else f'<dimension ref="{dimension}"/>'.encode()
    )
    parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(
        old, replacement
    )
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)
    book = engine.load_workbook(path, read_only=True)
    sheet = book.active
    if dimension:
        assert list(sheet.values) == [(1,)]
    else:
        assert sheet.max_row is None and sheet.max_column is None
        with pytest.raises(ValueError):
            sheet.calculate_dimension()
    sheet.reset_dimensions()
    assert sheet.calculate_dimension(force=True) == "A1:I5"
    assert len(list(sheet.values)) == 5
    book.close()


def test_read_only_independent_iterators_early_close_and_other_sheet(engine, tmp_path):
    path = tmp_path / "independent.xlsx"
    source(path)
    book = engine.load_workbook(path, read_only=True)
    first = book.active.iter_rows(values_only=True)
    second = book.active.iter_rows(min_row=3, values_only=True)
    assert next(first)[0] == 1
    assert next(second)[1] == "middle"
    assert list(book["Other"].values) == [(42,)]
    first.close()
    assert next(second)[0] is None
    second.close()
    book.close()


@pytest.mark.skipif(os.name == "nt", reason="POSIX removal of an open input")
def test_read_only_retains_the_original_source_handle(engine, tmp_path):
    path = tmp_path / "retained.xlsx"
    source(path)
    book = engine.load_workbook(path, read_only=True)
    path.unlink()
    assert next(book.active.values)[0] == 1
    book.close()


@pytest.mark.parametrize("iso", [False, True])
def test_write_only_interleaved_sheets_temporal_formulas_empty_rows_and_single_save(
    engine, iso, tmp_path
):
    book = engine.Workbook(write_only=True, iso_dates=iso)
    assert book.write_only and not book.read_only
    assert book.active is None and book.sheetnames == []
    first = book.create_sheet("First")
    second = book.create_sheet("Second")
    first.append([1, None, "=A1+1", datetime(2024, 2, 29, 12, 3, 4, 123000)])
    second.append([True, " text ", "#N/A"])
    first.append([])
    first.append(value for value in [2, time(12, 3, 4), timedelta(days=2, seconds=3)])
    second.append([42])
    first.title = "Renamed"
    second.close()
    book.active = 1
    path = tmp_path / "write.xlsx"
    book.save(path)
    with pytest.raises(Exception):
        book.save(tmp_path / "repeat.xlsx")
    with pytest.raises(Exception):
        first.append([3])
    book.close()
    verified = openpyxl.load_workbook(path)
    assert (
        verified.sheetnames == ["Renamed", "Second"]
        and verified.active.title == "Second"
    )
    assert verified["Renamed"]["D1"].value == datetime(2024, 2, 29, 12, 3, 4, 123000)
    assert verified["Renamed"]["C1"].value == "=A1+1"
    assert verified["Renamed"]["A2"].value is None
    assert verified["Renamed"]["B3"].value == time(12, 3, 4)
    assert verified["Renamed"]["C3"].value == timedelta(days=2, seconds=3)
    assert list(verified["Second"].values) == [
        (True, " text ", "#N/A"),
        (42, None, None),
    ]
    verified.close()


def test_write_only_cell_values_and_epoch(engine, tmp_path):
    from importlib import import_module

    cell_type = import_module(engine.__name__ + ".cell.cell").WriteOnlyCell
    book = engine.Workbook(write_only=True)
    book.epoch = datetime(1904, 1, 1)
    sheet = book.create_sheet()
    sheet.append(
        [
            cell_type(sheet, "hello"),
            cell_type(sheet, datetime(2020, 1, 2)),
            cell_type(sheet, "=1+2"),
        ]
    )
    path = tmp_path / "cells.xlsx"
    book.save(path)
    book.close()
    verified = openpyxl.load_workbook(path)
    assert verified.epoch == datetime(1904, 1, 1)
    assert next(verified.active.values) == ("hello", datetime(2020, 1, 2), "=1+2")
    verified.close()


def test_streaming_close_and_failed_save_remove_all_owned_files(tmp_path):
    spool = tmp_path / "spools"
    spool.mkdir()
    book = crabxl.Workbook(write_only=True, temp_directory=spool)
    first, second = book.create_sheet("First"), book.create_sheet("Second")
    first.append([1])
    second.append([2])
    assert len(list(spool.iterdir())) == 2
    book.close()
    book.close()
    assert list(spool.iterdir()) == []
    book = crabxl.Workbook(write_only=True, temp_directory=spool)
    book.create_sheet().append([1])
    target = tmp_path / "directory"
    target.mkdir()
    with pytest.raises(OSError):
        book.save(target)
    assert list(spool.iterdir()) == []
    assert target.is_dir()
    book.close()
    book = crabxl.Workbook(write_only=True, temp_directory=spool)
    book.create_sheet().append([1])
    with pytest.raises(OSError):
        book.save(tmp_path / "missing" / "output.xlsx")
    assert list(spool.iterdir()) == []
    book.close()


def test_read_only_close_cancels_live_iterator_and_modes_reject_unsupported_features(
    tmp_path,
):
    path = tmp_path / "close.xlsx"
    source(path)
    book = crabxl.load_workbook(path, read_only=True)
    rows = book.active.values
    next(rows)
    book.close()
    with pytest.raises(ValueError, match="closed"):
        next(rows)
    assert list(book._streams) == []
    from crabxl.cell.cell import WriteOnlyCell

    with pytest.raises(NotImplementedError):
        WriteOnlyCell().number_format = "0.00"


def test_read_only_selected_prefix_skips_invalid_unread_tail(tmp_path):
    path = tmp_path / "prefix.xlsx"
    source(path)
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/worksheets/sheet1.xml"] = (
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<dimension ref="A1:A2"/><sheetData>'
        '<row r="1"><c r="A1"><v>7</v></c></row>'
        '<row r="2"><c r="A2"><v>invalid</v></c></row>'
        "</sheetData></worksheet>"
    ).encode()
    with zipfile.ZipFile(path, "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    book = crabxl.load_workbook(path, read_only=True)
    assert list(book.active.iter_rows(max_row=1, values_only=True)) == [(7,)]
    with pytest.raises(ValueError):
        list(book.active.values)
    book.close()


def test_write_only_rejects_large_rows_before_consuming_unbounded_generator(tmp_path):
    book = crabxl.Workbook(write_only=True, max_memory_bytes=65536)
    sheet = book.create_sheet()
    consumed = 0

    def values():
        nonlocal consumed
        for _ in range(16384):
            consumed += 1
            yield "x" * 32767

    with pytest.raises(MemoryError):
        sheet.append(values())
    assert consumed <= 3
    sheet.append([7])
    path = tmp_path / "recovered.xlsx"
    book.save(path)
    check = openpyxl.load_workbook(path, read_only=True)
    assert list(check.active.values) == [(7,)]
    check.close()
    book.close()


def test_read_only_shared_strings_spill_and_partial_close_cleanup(
    tmp_path, monkeypatch, shared_strings_source
):
    path = tmp_path / "strings.xlsx"
    source(path)
    values = [f"value-{index}-" + "x" * 500 for index in range(500)]
    shared_strings_source(path, values)
    temporary = tmp_path / "sst-temporary"
    temporary.mkdir()
    monkeypatch.setenv("TMPDIR", str(temporary))
    monkeypatch.setenv("TEMP", str(temporary))
    monkeypatch.setenv("TMP", str(temporary))
    book = crabxl.load_workbook(path, read_only=True, max_memory_bytes=65536)
    rows = book.active.values
    assert next(rows) == (values[0],)

    # SST files are anonymous/delete-on-close; Linux exposes their live handles.
    def owned_handles():
        handles = []
        descriptors = Path("/proc/self/fd")
        if descriptors.is_dir():
            for descriptor in descriptors.iterdir():
                try:
                    if str(temporary) in os.readlink(descriptor):
                        handles.append(descriptor.name)
                except OSError:
                    pass
        return handles

    if Path("/proc/self/fd").is_dir():
        assert len(owned_handles()) == 2
    rows.close()
    assert owned_handles() == []
    assert list(temporary.iterdir()) == []
    assert list(book.active.values) == [(value,) for value in values]
    assert list(temporary.iterdir()) == []
    book.close()


def test_read_only_values_preserve_structured_formula_objects(tmp_path, engine):
    from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

    path = tmp_path / "structured.xlsx"
    book = openpyxl.Workbook()
    book.active["A1"] = ArrayFormula(ref="A1:A2", text="=SUM(B1:B2)")
    book.active["C1"] = DataTableFormula(ref="C1:C2", r1="B1", dt2D=True)
    book.active["B1"], book.active["B2"] = 1, 2
    book.save(path)
    book.close()
    book = engine.load_workbook(path, read_only=True)
    row = next(book.active.iter_rows(max_row=1, values_only=True))
    assert vars(row[0]) == {"ref": "A1:A2", "text": "=SUM(B1:B2)"}
    assert row[2].ref == "C1:C2"
    assert row[2].r1 == "B1"
    assert bool(row[2].dt2D)
    cell = book.active.cell(1, 1)
    assert cell.data_type == "f"
    assert vars(cell.value) == vars(row[0])
    book.close()
