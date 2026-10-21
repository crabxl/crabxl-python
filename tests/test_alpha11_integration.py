"""Consolidated A11 declaration, live-object and structural acceptance."""

import importlib
from copy import copy

import crabxl
import openpyxl
import pytest


@pytest.mark.parametrize("engine", [openpyxl, crabxl])
@pytest.mark.parametrize("loaded", [False, True])
def test_raw_merges_keep_realized_cells_and_live_coordinates(engine, loaded, tmp_path):
    ranges = importlib.import_module(engine.__name__ + ".worksheet.cell_range")
    book = engine.Workbook()
    sheet = book.active
    sheet["A1"] = "anchor"
    sheet.merge_cells("A1:C3")
    if loaded:
        source = tmp_path / "raw-source.xlsx"
        book.save(source)
        book.close()
        book = engine.load_workbook(source)
        sheet = book.active
    view = next(iter(sheet.merged_cells.ranges))
    view.shift(row_shift=2, col_shift=4)
    assert view.coord == "E3:G5"
    assert str(next(iter(sheet.merged_cells.ranges))) == "E3:G5"
    assert type(sheet["B2"]).__name__ == "MergedCell"
    assert type(sheet["F4"]).__name__ == "Cell"
    sheet["J2"] = "retained"
    sheet.merged_cells.ranges.add(ranges.CellRange("J1:K2"))
    assert sheet["J2"].value == "retained"
    sheet.merged_cells.remove("J1:K2")
    sheet.merged_cells.ranges = {view, ranges.CellRange("M1:N2")}
    view.shift(row_shift=1)
    assert "E4:G6" in {str(value) for value in sheet.merged_cells.ranges}
    with pytest.raises(ValueError):
        view.shift(row_shift=-100)
    assert view.coord == "E4:G6"
    detached = copy(sheet.merged_cells)
    sheet.merged_cells.ranges.clear()
    assert len(sheet.merged_cells.ranges) == 0
    assert len(detached.ranges) == 2
    assert type(sheet["B2"]).__name__ == "MergedCell"
    sheet.merged_cells.ranges = detached.ranges
    for suffix in (1, 2):
        output = tmp_path / f"raw-{suffix}.xlsx"
        book.save(output)
        reopened = openpyxl.load_workbook(output)
        assert {str(value) for value in reopened.active.merged_cells.ranges} == {
            "E4:G6",
            "M1:N2",
        }
        assert reopened.active["J2"].value == "retained"
        reopened.close()
    book.close()


@pytest.mark.parametrize("engine", [openpyxl, crabxl])
@pytest.mark.parametrize("loaded", [False, True])
def test_structural_merge_appearance_and_hyperlink_owners(engine, loaded, tmp_path):
    book = engine.Workbook()
    sheet = book.active
    sheet["A2"] = "anchor"
    sheet["A2"].border = engine.styles.Border(
        top=engine.styles.Side(style="thin"),
        right=engine.styles.Side(style="double"),
    )
    sheet.merge_cells("A2:C4")
    sheet["E2"].hyperlink = "https://example.org/old"
    if loaded:
        source = tmp_path / "structure-source.xlsx"
        book.save(source)
        book.close()
        book = engine.load_workbook(source)
        sheet = book.active
    link = sheet["E2"].hyperlink
    sheet.insert_rows(3)
    assert str(sheet.merged_cells) == "A2:C4"
    assert type(sheet["B3"]).__name__ == "Cell"
    assert type(sheet["B5"]).__name__ == "MergedCell"
    assert sheet["C5"].border.right.style == "double"
    sheet.move_range("A2:F5", rows=3, cols=1)
    assert sheet["B5"].value == "anchor"
    assert type(sheet["C8"]).__name__ == "MergedCell"
    assert sheet["D8"].border.right.style == "double"
    assert sheet["F5"].hyperlink is link
    assert link.ref == "E2"
    link.target = "https://example.org/new"
    assert sheet["F5"].hyperlink.target == "https://example.org/new"
    sheet.delete_cols(1)
    assert sheet["E5"].hyperlink is link and link.ref == "E2"
    for suffix in (1, 2):
        output = tmp_path / f"structure-{suffix}.xlsx"
        book.save(output)
        reopened = openpyxl.load_workbook(output)
        assert reopened.active["E2"].hyperlink.target == "https://example.org/new"
        assert reopened.active["E5"].value == "https://example.org/old"
        assert reopened.active["C8"].border.right.style == "double"
        reopened.close()
    book.close()


@pytest.mark.parametrize("engine", [openpyxl, crabxl])
def test_rich_copy_shared_runs_and_optimization_do_not_bind_new_list(engine):
    rich = importlib.import_module(engine.__name__ + ".cell.rich_text")
    font = importlib.import_module(engine.__name__ + ".cell.text").InlineFont
    book = engine.Workbook()
    block = rich.TextBlock(font(b=True), "run")
    value = rich.CellRichText(block, "a", "b", "")
    book.active["A1"] = value
    duplicate = copy(value)
    assert duplicate[0] is block
    duplicate.append("copy")
    assert str(book.active["A1"].value) == "runab"
    block.text = "shared"
    assert str(book.active["A1"].value) == "sharedab"
    assert str(duplicate) == "sharedabcopy"
    value._opt()
    assert value.as_list() == ["shared", "ab"]
    book.close()


@pytest.mark.parametrize("engine", [openpyxl, crabxl])
@pytest.mark.parametrize("mode", ["owned", "loaded", "write_only"])
def test_named_theme_dimension_styles_grouping_and_repeat_output(
    engine, mode, tmp_path
):
    book = engine.Workbook(write_only=mode == "write_only")
    sheet = book.create_sheet("Data") if mode == "write_only" else book.active
    style = engine.styles.NamedStyle(
        name="Heading",
        font=engine.styles.Font(b=True, color="FF123456"),
        number_format="0.00",
        protection=engine.styles.Protection(locked=False),
    )
    book.add_named_style(style)
    sheet.row_dimensions[2].height = 27
    sheet.row_dimensions[2].font = engine.styles.Font(i=True)
    sheet.column_dimensions["B"].width = 24
    sheet.column_dimensions["B"].number_format = "0.000"
    sheet.column_dimensions.group("C", "E", outline_level=2, hidden=True)
    sheet.row_dimensions.group(4, 6, outline_level=1, hidden=True)
    if mode == "write_only":
        sheet.row_dimensions[99].height = 50
        cell = engine.cell.WriteOnlyCell(sheet, value=12.5)
        cell.style = "Heading"
        sheet.append([cell])
        for _ in range(5):
            sheet.append([])
    else:
        sheet["A1"] = 12.5
        sheet["A1"].style = "Heading"
    if mode == "loaded":
        source = tmp_path / "named-source.xlsx"
        book.save(source)
        book.close()
        book = engine.load_workbook(source)
        sheet = book.active
        assert "Heading" in book.named_styles
        sheet.row_dimensions[2].height = 31
        sheet["B1"] = 7.5
        sheet["B1"].style = "Heading"
        book.loaded_theme = book.loaded_theme.replace(
            b'name="Office Theme"', b'name="A11 Theme"'
        )
    for suffix in range(1 if mode == "write_only" else 2):
        output = tmp_path / f"named-{suffix}.xlsx"
        book.save(output)
        reference = openpyxl.load_workbook(output)
        target = reference.active
        assert target["A1"].font.b and target["A1"].font.color.rgb == "FF123456"
        assert target["A1"].number_format == "0.00"
        assert target["A1"].style == "Heading"
        assert not target["A1"].protection.locked
        assert target.row_dimensions[2].height == (31 if mode == "loaded" else 27)
        assert target.row_dimensions[2].font.i
        assert target.column_dimensions["B"].width == 24
        assert target.column_dimensions["B"].number_format == "0.000"
        assert (
            target.column_dimensions["C"].min == 3
            and target.column_dimensions["C"].max == 5
        )
        assert (
            target.column_dimensions["C"].hidden
            and target.column_dimensions["C"].outlineLevel == 2
        )
        assert all(target.row_dimensions[row].hidden for row in (4, 5, 6))
        if mode == "write_only":
            assert 99 not in target.row_dimensions
        if mode == "loaded":
            assert b"A11 Theme" in reference.loaded_theme
            assert target["B1"].style == "Heading"
        reference.close()
    book.close()


@pytest.mark.parametrize("engine", [openpyxl, crabxl])
@pytest.mark.parametrize("loaded", [False, True])
def test_merging_drops_nonanchor_links_but_retains_detached_alias(
    engine, loaded, tmp_path
):
    book = engine.Workbook()
    sheet = book.active
    sheet["A1"].hyperlink = "https://example.org/anchor"
    sheet["B2"].hyperlink = "https://example.org/interior"
    if loaded:
        source = tmp_path / "merge-links-source.xlsx"
        book.save(source)
        book.close()
        book = engine.load_workbook(source)
        sheet = book.active
    alias = sheet["B2"]
    sheet.merge_cells("A1:C3")
    assert sheet["B2"].hyperlink is None
    assert alias.hyperlink.target == "https://example.org/interior"
    assert sheet["A1"].hyperlink.target == "https://example.org/anchor"
    for suffix in (1, 2):
        output = tmp_path / f"merge-links-{suffix}.xlsx"
        book.save(output)
        reference = openpyxl.load_workbook(output)
        assert reference.active["A1"].hyperlink.target == "https://example.org/anchor"
        assert reference.active["B2"].hyperlink is None
        reference.close()
    book.close()


@pytest.mark.parametrize("engine", [openpyxl, crabxl])
def test_write_only_shared_hyperlink_tracks_closed_sheet_and_saved_ids(
    engine, tmp_path
):
    links = importlib.import_module(engine.__name__ + ".worksheet.hyperlink")
    book = engine.Workbook(write_only=True)
    first = book.create_sheet("First")
    second = book.create_sheet("Second")
    link = links.Hyperlink(ref="Z9", target="https://example.org/original")
    left = engine.cell.WriteOnlyCell(first, value="first label")
    left.hyperlink = link
    first.append([left])
    first.close()
    assert link.id == "rId1"
    right = engine.cell.WriteOnlyCell(second, value="second label")
    right.hyperlink = link
    second.append([None, right])
    link.target = "https://example.org/changed"
    link.tooltip = "shared after closing first sheet"
    output = tmp_path / "stream-links.xlsx"
    book.save(output)
    assert link.ref == "B1" and link.id == "rId1"
    reopened = openpyxl.load_workbook(output)
    assert reopened["First"]["A1"].hyperlink.target == "https://example.org/original"
    assert reopened["First"]["A1"].hyperlink.tooltip is None
    assert reopened["First"]["B1"].hyperlink is None
    assert reopened["Second"]["A1"].hyperlink is None
    assert reopened["Second"]["B1"].hyperlink.target == link.target
    assert reopened["Second"]["B1"].hyperlink.tooltip == link.tooltip
    assert reopened["First"]["A1"].value == "first label"
    assert reopened["Second"]["B1"].value == "second label"
    reopened.close()
    book.close()


def test_shared_hyperlink_budget_failure_rolls_back_all_live_owners():
    from crabxl.worksheet.hyperlink import Hyperlink

    large = crabxl.Workbook(max_memory_bytes=256 * 1024)
    small = crabxl.Workbook(max_memory_bytes=32 * 1024)
    link = Hyperlink(ref="A1", target="https://example.org/original")
    large.active["A1"].hyperlink = link
    small.active["A1"].hyperlink = link
    with pytest.raises(MemoryError):
        link.target = "x" * (64 * 1024)
    assert link.target == "https://example.org/original"
    assert large.active["A1"].hyperlink.target == link.target
    assert small.active["A1"].hyperlink.target == link.target
    link.tooltip = "retry"
    assert (
        large.active["A1"].hyperlink.tooltip
        == small.active["A1"].hyperlink.tooltip
        == "retry"
    )
    large.close()
    small.close()
