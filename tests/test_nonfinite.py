"""Public nonfinite behavior uses canonical Rust serialization, without fallback."""

import math
import zipfile

import crabxl
import openpyxl
import pytest


@pytest.fixture(params=[openpyxl, crabxl], ids=["openpyxl", "crabxl"])
def engine(request):
    return request.param


@pytest.mark.parametrize(
    "value",
    [float("nan"), float("inf"), -float("inf")],
    ids=["nan", "inf", "negative-inf"],
)
def test_nonfinite_owned_and_original_editor_values(engine, value, tmp_path):
    book = engine.Workbook()
    book.active["A1"] = value
    actual = book.active["A1"].value
    assert math.isnan(actual) if math.isnan(value) else actual == value
    path = tmp_path / "owned.xlsx"
    book.save(path)
    book.close()
    loaded = openpyxl.load_workbook(path)
    assert loaded.active["A1"].value is None
    loaded.close()
    loaded = engine.load_workbook(path)
    loaded.active["A1"] = value
    edited = tmp_path / "edited.xlsx"
    loaded.save(edited)
    loaded.close()
    checked = openpyxl.load_workbook(edited)
    assert checked.active["A1"].value is None
    checked.close()


@pytest.mark.parametrize("cached", [False, True])
def test_scientific_overflow_and_formula_cache(engine, cached, tmp_path):
    path = tmp_path / "source.xlsx"
    book = openpyxl.Workbook()
    book.save(path)
    book.close()
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/worksheets/sheet1.xml"] = (
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1"><v>1e999</v></c><c r="B1"><v>-1e999</v></c><c r="C1"><f>1</f><v>1e999</v></c></row></sheetData></worksheet>'
    )
    with zipfile.ZipFile(path, "w") as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)
    book = engine.load_workbook(path, data_only=cached)
    assert book.active["A1"].value == float("inf")
    assert book.active["B1"].value == -float("inf")
    assert book.active["C1"].value == (float("inf") if cached else "=1")
    book.close()
