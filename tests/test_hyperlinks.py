"""Public point lifecycle and source graph interoperability."""

from copy import copy
from xml.etree.ElementTree import tostring

import openpyxl
import pytest
from crabxl import Workbook, load_workbook
from crabxl.worksheet.hyperlink import Hyperlink, HyperlinkList


def test_live_point_hyperlinks_preserve_values_and_source_copy_graphs(tmp_path):
    book = Workbook()
    sheet = book.active
    link = Hyperlink(
        ref="Z8", target="https://example.org/old?a=1&b=2#part", tooltip="tip"
    )
    sheet["A1"].hyperlink = link
    assert sheet["A1"].hyperlink is link
    assert link.ref == "A1"
    assert sheet["A1"].value == link.target
    link.target = "../new.xlsx#Sheet!A1"
    assert sheet["A1"].value == "https://example.org/old?a=1&b=2#part"
    sheet["B2"].hyperlink = Hyperlink(ref="B2", location="'Local Sheet'!A1")
    assert sheet["B2"].value == "'Local Sheet'!A1"
    sheet["C3"] = "kept"
    sheet["C3"].hyperlink = ""
    assert sheet["C3"].hyperlink.target == ""
    sheet["C3"].hyperlink = None
    assert sheet["C3"].value == "kept"
    path = tmp_path / "points.xlsx"
    book.save(path)
    other = openpyxl.load_workbook(path)
    assert other.active["A1"].hyperlink.target == "../new.xlsx#Sheet!A1"
    assert other.active["A1"].hyperlink.tooltip == "tip"
    other.close()

    loaded = load_workbook(path)
    original = loaded.active
    current = original["A1"].hyperlink
    assert original["A1"].hyperlink is current
    current.target = "https://example.org/changed"
    current.tooltip = "changed tip"
    duplicate = loaded.copy_worksheet(original)
    duplicate["A1"].hyperlink.target = "https://example.org/copied"
    del original["B2"]
    assert original["B2"].hyperlink is None
    for index in range(2):
        output = tmp_path / f"edited-{index}.xlsx"
        loaded.save(output)
        other = openpyxl.load_workbook(output)
        assert (
            other.worksheets[0]["A1"].hyperlink.target == "https://example.org/changed"
        )
        assert other.worksheets[0]["A1"].hyperlink.tooltip == "changed tip"
        assert other.worksheets[0]["B2"].hyperlink is None
        assert (
            other.worksheets[1]["A1"].hyperlink.target == "https://example.org/copied"
        )
        assert other.worksheets[0]["A1"].value == "https://example.org/old?a=1&b=2#part"
        other.close()
    with pytest.raises(NotImplementedError):
        current.ref = "D4"
    assert current.ref == "A1"
    with pytest.raises(NotImplementedError):
        original["D4"].hyperlink = current
    assert original["D4"].hyperlink is None
    loaded.close()


def test_hyperlink_xml_public_fields_match_reference():
    fields = dict(
        ref="A1",
        target="https://example.org/",
        location="B2",
        tooltip="tip",
        display="label",
        id="rId7",
    )
    actual = Hyperlink(**fields)
    expected = openpyxl.worksheet.hyperlink.Hyperlink(**fields)
    assert actual.to_tree().attrib == expected.to_tree().attrib
    assert "target" not in actual.to_tree().attrib
    restored = Hyperlink.from_tree(actual.to_tree())
    assert restored.target is None
    assert restored.id == "rId7"
    assert copy(actual) == actual
    changed = copy(actual)
    changed.target = "different target"
    assert changed == actual  # Reference equality compares serialized attributes.
    assert dict(actual) == dict(expected)
    collection = HyperlinkList(hyperlink=[actual])
    reference = openpyxl.worksheet.hyperlink.HyperlinkList(hyperlink=[expected])
    assert tostring(collection.to_tree()) == tostring(reference.to_tree())
    assert HyperlinkList.from_tree(collection.to_tree()).hyperlink[0].id == "rId7"
    assert tostring(restored.to_tree()) == tostring(actual.to_tree())
    with pytest.raises(TypeError):
        Hyperlink()
