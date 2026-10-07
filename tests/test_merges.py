"""Merged geometry, shared appearance and preserving saves through public APIs."""

import importlib
import zipfile

import crabxl
import openpyxl
import pytest


@pytest.mark.parametrize("engine", [openpyxl, crabxl], ids=["openpyxl", "crabxl"])
def test_merged_values_borders_aliases_and_preserving_repeat_saves(engine, tmp_path):
    for loaded in [False, True]:
        book = engine.Workbook()
        sheet = book.active
        sheet["A1"] = "anchor"
        sheet["B2"] = "discarded"
        sheet["A1"].border = engine.styles.Border(
            left=engine.styles.Side(style="thin"),
            top=engine.styles.Side(style="dashed"),
        )
        sheet["A1"].protection = engine.styles.Protection(locked=False, hidden=True)
        sheet["C3"].border = engine.styles.Border(
            right=engine.styles.Side(style="double"),
            bottom=engine.styles.Side(style="thick"),
        )
        if loaded:
            source = tmp_path / f"{engine.__name__}-source.xlsx"
            book.save(source)
            book.close()
            # A valid source may declare merges without pre-normalized edge styles.
            with zipfile.ZipFile(source) as archive:
                parts = {name: archive.read(name) for name in archive.namelist()}
            parts["xl/worksheets/sheet1.xml"] = parts[
                "xl/worksheets/sheet1.xml"
            ].replace(
                b"</sheetData>",
                b'</sheetData><mergeCells count="1"><mergeCell ref="A1:C3"/></mergeCells>',
            )
            with zipfile.ZipFile(
                source, "w", compression=zipfile.ZIP_DEFLATED
            ) as archive:
                for name, value in parts.items():
                    archive.writestr(name, value)
            book = engine.load_workbook(source)
            sheet = book.active
            assert sheet["B2"].value is None
            if engine is crabxl:
                untouched = tmp_path / "untouched.xlsx"
                book.save(untouched)
                with zipfile.ZipFile(untouched) as archive:
                    assert {
                        name: archive.read(name) for name in archive.namelist()
                    } == parts
        else:
            alias = sheet["B2"]
            sheet.merge_cells("A1:C3")
            assert alias.value == "discarded"
            alias.value = "detached"
            assert sheet["B2"].value is None
        merged_type = importlib.import_module(engine.__name__ + ".cell.cell").MergedCell
        assert isinstance(sheet["B2"], merged_type)
        with pytest.raises(AttributeError):
            sheet["B2"].value = 42
        assert sheet["A1"].value == "anchor"
        assert sheet.max_row == 3 and sheet.max_column == 3
        assert "B2" in sheet.merged_cells
        assert {str(value) for value in sheet.merged_cells.ranges} == {"A1:C3"}
        assert len(sheet.merged_cells.ranges) == 1
        if engine is crabxl:
            from crabxl.worksheet.cell_range import CellRange

            with pytest.raises(NotImplementedError):
                sheet.merged_cells.ranges.add(CellRange("J1:K2"))
            with pytest.raises(NotImplementedError):
                next(iter(sheet.merged_cells.ranges)).shift(row_shift=1)
            assert {str(value) for value in sheet.merged_cells.ranges} == {"A1:C3"}

        assert sheet["A1"].border.right.style == "double"
        assert sheet["A1"].border.bottom.style == "thick"
        for coord, side, expected in [
            ("A2", "left", "thin"),
            ("B1", "top", "dashed"),
            ("C2", "right", "double"),
            ("B3", "bottom", "thick"),
        ]:
            cell = sheet[coord]
            assert getattr(cell.border, side).style == expected
            assert not cell.protection.locked and cell.protection.hidden
        sheet["A1"] = "changed"
        sheet.merge_cells("E5:F6")
        for suffix in [1, 2]:
            path = tmp_path / f"{engine.__name__}-{loaded}-{suffix}.xlsx"
            book.save(path)
            reopened = openpyxl.load_workbook(path)
            target = reopened.active
            assert target["A1"].value == "changed"
            assert target["A1"].border.right.style == "double"
            assert target["B3"].border.bottom.style == "thick"
            assert {str(value) for value in target.merged_cells.ranges} == {
                "A1:C3",
                "E5:F6",
            }
            reopened.close()
        sheet.unmerge_cells("A1:C3")
        assert sheet["A1"].value == "changed" and sheet["B2"].value is None
        assert not isinstance(sheet["B2"], merged_type)
        sheet["B2"] = "restored"
        assert sheet["B2"].value == "restored"
        book.close()
