"""Compression controls preserve values and protect targets on invalid options."""

from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

import crabxl
import openpyxl
import pytest


def test_compression_levels_cover_creation_streaming_and_loaded_edits(tmp_path):
    values = ["Unicode 中文 & < >", True, 2**80, "=1+1"]
    source = tmp_path / "source.xlsx"
    book = crabxl.Workbook()
    book.active.append(values)
    book.save(source)
    for mode in ["ordinary", "write_only", "loaded"]:
        expected_xml = None
        book = None
        for level in [None, 0, 1, 3, 6, 9]:
            if mode == "loaded" and book is None:
                book = crabxl.load_workbook(source)
                book.active["A1"] = values[0] + " edited"
            elif book is None or mode == "write_only":
                book = crabxl.Workbook(write_only=mode == "write_only")
                sheet = (
                    book.create_sheet("Sheet") if mode == "write_only" else book.active
                )
                sheet.append(values)
            path = tmp_path / f"{mode}-{level}.xlsx"
            book.save(path, compression_level=level)
            if mode == "write_only":
                book.close()
            with ZipFile(path) as archive:
                xml = archive.read("xl/worksheets/sheet1.xml")
                if expected_xml is not None:
                    assert xml == expected_xml, (mode, level)
                expected_xml = xml
                assert archive.getinfo("xl/worksheets/sheet1.xml").compress_type == (
                    ZIP_STORED if level == 0 else ZIP_DEFLATED
                )
            check = openpyxl.load_workbook(path, read_only=True)
            expected = list(values)
            if mode == "loaded":
                expected[0] += " edited"
            assert next(check.active.values) == tuple(expected), (mode, level)
            check.close()
        book.close()


def test_invalid_compression_does_not_consume_stream_or_replace_target(tmp_path):
    path = tmp_path / "target.xlsx"
    for write_only in [False, True]:
        book = crabxl.Workbook(write_only=write_only)
        sheet = book.create_sheet("Sheet") if write_only else book.active
        sheet.append([7])
        path.write_bytes(b"original target")
        for level, exception in [
            (-1, ValueError),
            (10, ValueError),
            (True, TypeError),
            (1.5, TypeError),
        ]:
            with pytest.raises(exception):
                book.save(path, compression_level=level)
            assert path.read_bytes() == b"original target"
        book.save(path, compression_level=1)
        book.close()
        check = crabxl.load_workbook(path, read_only=True)
        assert next(check.active.values) == (7,)
        check.close()
