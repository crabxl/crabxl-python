"""One public-API test body runs against both implementations."""

import warnings
import zipfile
from contextlib import nullcontext
from datetime import date, datetime, time, timedelta

import crabxl
import openpyxl
import pytest


@pytest.fixture(params=[openpyxl, crabxl], ids=["openpyxl", "crabxl"])
def engine(request):
    return request.param


def test_sheet_visibility_and_deferred_active_views(engine, tmp_path):
    cases = {
        2: [(-3, 0, 0), (-1, 0, 0), (1, 1, 0), (10, 10, 0)],
        3: [(-3, -3, 0), (-1, -1, 2), (1, 2, 2), (10, 10, 0)],
    }

    def active_title(book):
        return book.active.title if book.active is not None else None

    def selected_title(index, count):
        return ["A", "B", "C"][:count][index] if -count <= index < count else None

    for count, matrix in cases.items():
        for mode in ["owned", "loaded", "write_only"]:
            for requested, after, read_index in matrix:
                book = engine.Workbook(write_only=mode == "write_only")
                for index, title in enumerate(["A", "B", "C"][:count]):
                    sheet = (
                        book.active
                        if index == 0 and mode != "write_only"
                        else book.create_sheet(title)
                    )
                    sheet.title = title
                    sheet.append([index + 11])
                book["B"].sheet_state = "hidden"
                with pytest.raises(ValueError):
                    book.active = book["B"]
                if mode == "loaded":
                    source = tmp_path / f"source-{count}-{requested}.xlsx"
                    book.save(source)
                    book.close()
                    book = engine.load_workbook(source)
                book.active = requested
                assert active_title(book) == selected_title(requested, count)
                output = tmp_path / f"view-{count}-{mode}-{requested}.xlsx"
                for _ in range(1 if mode == "write_only" else 2):
                    book.save(output)
                    assert active_title(book) == selected_title(after, count)
                    checked = openpyxl.load_workbook(output, read_only=True)
                    assert active_title(checked) == ["A", "B", "C"][read_index]
                    assert checked["B"].sheet_state == "hidden"
                    assert [list(sheet.values) for sheet in checked] == [
                        [(index + 11,)] for index in range(count)
                    ]
                    checked.close()
                read = engine.load_workbook(output, read_only=True)
                assert read["B"].sheet_state == "hidden"
                # Read-only metadata changes affect the view, never source output.
                read["B"].sheet_state = "veryHidden"
                assert read["B"].sheet_state == "veryHidden"
                read.close()
                book.close()

    for loaded in [False, True]:
        book = engine.Workbook()
        book.active.title = "A"
        book.create_sheet("B")
        book["A"].append([11])
        book["B"].append([22])
        if loaded:
            source = tmp_path / "hidden-source.xlsx"
            book.save(source)
            book.close()
            book = engine.load_workbook(source)
        book["A"].sheet_state = "hidden"
        book["B"].sheet_state = "veryHidden"
        output = tmp_path / f"hidden-{loaded}.xlsx"
        with pytest.raises(IndexError):
            book.save(output)
        book["B"].sheet_state = "visible"
        book.save(output)
        assert active_title(book) == "B"
        read = openpyxl.load_workbook(output)
        assert read["A"].sheet_state == "hidden"
        assert read["B"].sheet_state == "visible"
        assert read["B"]["A1"].value == 22
        read.close()
        book.close()

    book = engine.Workbook()
    book.active.sheet_state = "hidden"
    copied = book.copy_worksheet(book.active)
    assert copied.sheet_state == "visible"
    book.remove(copied)
    with pytest.raises(ValueError):
        book.save(tmp_path / "only-hidden.xlsx")
    book.close()


def test_scalar_types_formulas_dimensions_and_live_views(engine):
    workbook = engine.Workbook()
    sheet = workbook.active
    values = [None, True, False, 5, 1.25, " whitespace ", "#DIV/0!", "=1+2"]
    sheet.append(values)
    assert list(sheet.values) == [tuple(values)]
    assert [sheet.cell(1, column).data_type for column in range(1, 9)] == [
        "n",
        "b",
        "b",
        "n",
        "n",
        "s",
        "e",
        "f",
    ]
    cell = sheet["D1"]
    assert sheet["D1"] is cell
    sheet.insert_rows(1, 2)
    assert cell.coordinate == "D3" and cell.value == 5
    sheet.delete_cols(1, 2)
    assert cell.coordinate == "B3" and sheet["B3"] is cell
    sheet.move_range("B3:C3", rows=1, cols=2)
    assert cell.coordinate == "D4" and sheet["D4"] is cell
    del sheet["D4"]
    assert cell.value == 5 and sheet["D4"].value is None
    cell.value = 42
    assert sheet["D4"].value is None


def test_new_save_public_readback_and_repeat(engine, tmp_path):
    workbook = engine.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(
        [
            True,
            5,
            1.25,
            " text ",
            "#N/A",
            "=SUM(B1:C1)",
            datetime(2020, 2, 29, 12, 3, 4, 123000),
            time(12, 3, 4),
            timedelta(days=2, seconds=3),
        ]
    )
    workbook.create_sheet("Other").append(["second"])
    path = tmp_path / "new.xlsx"
    workbook.save(path)
    book = openpyxl.load_workbook(path)
    assert book.sheetnames == ["Data", "Other"]
    row = tuple(book["Data"].values)[0]
    assert row[:6] == (True, 5, 1.25, " text ", "#N/A", "=SUM(B1:C1)")
    assert row[6:] == (
        datetime(2020, 2, 29, 12, 3, 4, 123000),
        time(12, 3, 4),
        timedelta(days=2, seconds=3),
    )
    book.close()
    workbook.save(tmp_path / "repeat.xlsx")
    assert set(tmp_path.iterdir()) == {path, tmp_path / "repeat.xlsx"}


def test_loaded_numeric_edit_and_sparse_insertion(engine, tmp_path):
    from openpyxl.workbook.defined_name import DefinedName

    source = tmp_path / "source.xlsx"
    original = openpyxl.Workbook()
    original.active.title = "First"
    original.active.append([1, 2, "=A1+B1"])
    original.create_sheet("Second")["A1"] = "=First!A1"
    original.defined_names.add(DefinedName("Pick", attr_text="'First'!$A$1"))
    original.save(source)
    original.close()
    workbook = engine.load_workbook(source)
    second = workbook["Second"]
    second.title = "first"
    assert second.title == "first1"
    assert workbook["first1"] is second
    sheet = workbook.active
    alias = sheet["A1"]
    assert sheet["B1"].value == 2
    sheet.title = 'Renamed<&" \u65b0'
    assert workbook[sheet.title] is sheet
    assert alias.parent is sheet
    with pytest.raises(ValueError):
        sheet.title = "Invalid/Name"
    assert sheet.title == 'Renamed<&" \u65b0'
    sheet["A1"] = 5
    sheet["D5"] = "new"
    assert sheet["A1"].value == 5 and sheet["D5"].value == "new"
    assert alias.value == 5
    assert sheet.max_row == 5 and sheet.max_column == 4
    assert list(sheet.iter_rows(min_row=5, max_row=5, values_only=True)) == [
        (None, None, None, "new")
    ]
    sheet.append({"C": "appended", "D": "=A1+B1"})
    sheet.append([])
    sheet.append([8, True])
    assert sheet["C6"].value == "appended"
    assert sheet["D6"].value == "=A1+B1"
    assert sheet["A8"].value == 8 and sheet["B8"].value is True
    workbook.move_sheet(second, offset=-1)
    assert workbook.active is second
    assert workbook.worksheets == [second, sheet]
    assert alias.parent is sheet and alias.value == 5
    workbook.active = sheet
    output = tmp_path / "output.xlsx"
    workbook.save(output)
    workbook.save(tmp_path / "repeat.xlsx")
    workbook.close()
    verified = openpyxl.load_workbook(output)
    assert verified.sheetnames == ["first1", 'Renamed<&" \u65b0']
    assert verified["first1"]["A1"].value == "=First!A1"
    assert verified.defined_names["Pick"].attr_text == "'First'!$A$1"
    assert verified.active["A1"].value == 5 and verified.active["D5"].value == "new"
    assert verified.active["C1"].value == "=A1+B1"
    assert verified.active["C6"].value == "appended"
    assert verified.active["D6"].value == "=A1+B1"
    assert verified.active["A8"].value == 8 and verified.active["B8"].value is True
    verified.close()

    # Retain one shared loaded workflow for alias coordinates, structural edits,
    # translated formulas, later scalar edits/append, and repeat-save readback.
    workbook = engine.load_workbook(output)
    sheet = workbook.active
    alias = sheet["A1"]
    sheet.insert_rows(1, 2)
    assert alias.coordinate == "A3" and alias.value == 5
    sheet.delete_rows(1, 2)
    sheet.insert_cols(1, 2)
    assert alias.coordinate == "C1" and alias.value == 5
    sheet.delete_cols(1, 2)
    assert alias.coordinate == "A1" and alias.value == 5
    sheet.move_range("C1", rows=2, cols=2, translate=True)
    assert sheet["C1"].value is None
    assert sheet["E3"].value == "=C3+D3"
    sheet["D5"] = "after shift"
    sheet.append([9])
    assert sheet["A9"].value == 9
    for filename in ("shifted.xlsx", "shifted-repeat.xlsx"):
        workbook.save(tmp_path / filename)
        verified = openpyxl.load_workbook(tmp_path / filename)
        assert verified.active["A1"].value == 5
        assert verified.active["E3"].value == "=C3+D3"
        assert verified.active["D5"].value == "after shift"
        assert verified.active["A9"].value == 9
        assert verified["first1"]["A1"].value == "=First!A1"
        assert verified.defined_names["Pick"].attr_text == "'First'!$A$1"
        verified.close()
    workbook.close()


@pytest.mark.parametrize("loaded", [False, True])
def test_values_iteration_observes_edits_between_rows(engine, loaded, tmp_path):
    workbook = engine.Workbook()
    worksheet = workbook.active
    worksheet.append([1, "first", None])
    worksheet.append([2, "second", "=A2+1"])
    if loaded:
        path = tmp_path / "values.xlsx"
        workbook.save(path)
        workbook.close()
        workbook = engine.load_workbook(path)
        worksheet = workbook.active
    rows = worksheet.iter_rows(max_row=3, max_col=3, values_only=True)
    assert next(rows) == (1, "first", None)
    worksheet["B2"] = "changed"
    worksheet["C3"] = False
    assert next(rows) == (2, "changed", "=A2+1")
    assert next(rows) == (None, None, False)
    assert list(rows) == []
    assert list(
        worksheet.iter_rows(
            min_row=2, max_row=2, min_col=2, max_col=3, values_only=True
        )
    ) == [("changed", "=A2+1")]
    worksheet["F2"] = "tail"
    assert list(
        worksheet.iter_rows(
            min_row=2, max_row=2, min_col=4, max_col=7, values_only=True
        )
    ) == [(None, None, "tail", None)]
    assert list(
        worksheet.iter_rows(
            min_row=4, max_row=4, min_col=4, max_col=7, values_only=True
        )
    ) == [(None, None, None, None)]
    workbook.close()


def test_values_iteration_retains_live_structured_formula_identity(tmp_path):
    from crabxl.worksheet.formula import ArrayFormula

    path = tmp_path / "values-formula.xlsx"
    book = crabxl.Workbook()
    book.active["A1"] = ArrayFormula("A1:A2", "=SUM(B1:B2)")
    book.save(path)
    book.close()
    book = crabxl.load_workbook(path)
    cell = book.active["A1"]
    value = cell.value
    assert (
        next(book.active.iter_rows(max_row=1, max_col=1, values_only=True))[0] is value
    )
    value.text = "=SUM(C1:C2)"
    assert cell.value.text == "=SUM(C1:C2)"
    book.save(path)
    book.close()
    verified = openpyxl.load_workbook(path)
    assert verified.active["A1"].value.text == "=SUM(C1:C2)"
    verified.close()


def test_loaded_formula_data_only(engine, tmp_path):
    source = tmp_path / "formula.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["=1+2"])
    workbook.save(source)
    loaded = engine.load_workbook(source, data_only=True)
    assert loaded.active["A1"].value is None
    loaded.close()


def test_empty_append_cursor_and_dictionary_columns(engine):
    workbook = engine.Workbook()
    sheet = workbook.active
    sheet.append([])
    sheet.append({"C": 3})
    assert sheet["C2"].value == 3
    sheet.delete_rows(1, 2)
    sheet.append([5])
    assert sheet["A1"].value == 5


def test_loaded_style_image_comment_and_unknown_parts_are_preserved(tmp_path):
    from openpyxl.comments import Comment
    from openpyxl.drawing.image import Image as DrawingImage
    from openpyxl.styles import Font
    from openpyxl.workbook.defined_name import DefinedName
    from PIL import Image

    image_path = tmp_path / "image.png"
    Image.new("RGB", (8, 8), "red").save(image_path)
    source = tmp_path / "source.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = 1
    workbook.active["A1"].font = Font(bold=True)
    workbook.active["A1"].comment = Comment("comment", "author")
    workbook.active.add_image(DrawingImage(image_path), "D1")
    workbook.active.defined_names.add(DefinedName("Local", attr_text="Sheet!$A$1"))
    workbook.create_sheet("Other")
    workbook.save(source)
    with zipfile.ZipFile(source, "a") as archive:
        archive.writestr("custom/unknown.xml", '<x xmlns="urn:custom">keep</x>')
    loaded = crabxl.load_workbook(source)
    loaded.active.title = "Renamed"
    with pytest.raises(NotImplementedError, match="Local defined-name"):
        loaded.move_sheet(loaded.active, offset=1)
    assert loaded.sheetnames == ["Renamed", "Other"]
    cell = loaded.active["A1"]
    with pytest.raises(NotImplementedError, match="feature graphs"):
        loaded.active.insert_rows(1)
    with pytest.raises(NotImplementedError, match="feature graphs"):
        loaded.active.move_range("A1", rows=1)
    assert cell.coordinate == "A1" and cell.value == 1
    loaded.active["A1"] = 9
    for filename in ("first.xlsx", "second.xlsx"):
        target = tmp_path / filename
        loaded.save(target)
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(target) as saved:
            for name in ("xl/media/image1.png", "xl/styles.xml", "custom/unknown.xml"):
                assert original.read(name) == saved.read(name)
        verified = openpyxl.load_workbook(target)
        assert verified.active.title == "Renamed"
        assert verified.active.defined_names["Local"].attr_text == "Sheet!$A$1"
        assert verified.active["A1"].value == 9
        assert verified.active["A1"].font.bold
        assert verified.active["A1"].comment.text == "comment"
        verified.close()
    loaded.close()


def test_adapter_limits_exact_integers_closed_sources_and_unsupported_operations(
    tmp_path,
):
    workbook = crabxl.Workbook(max_memory_bytes=1000)
    sheet = workbook.active
    sheet["A1"] = 10**100
    assert sheet["A1"].value == 10**100
    with pytest.raises(MemoryError):
        sheet["A1"] = "x" * 2000
    assert sheet["A1"].value == 10**100
    with pytest.raises(AttributeError):
        sheet["A1"].font = object()
    with pytest.raises(AttributeError):
        sheet.freeze_panes = "A1"
    streaming = crabxl.Workbook(write_only=True)
    with pytest.raises(NotImplementedError):
        streaming.create_sheet().cell(1, 1)
    streaming.close()
    source = tmp_path / "source.xlsx"
    workbook.save(source)
    loaded = crabxl.load_workbook(source)
    loaded.close()
    with pytest.raises(ValueError, match="closed"):
        loaded.active["A1"].value
    assert len(list(tmp_path.iterdir())) == 1


def test_atomic_new_output_failure_preserves_existing_target(tmp_path):
    target = tmp_path / "target.xlsx"
    target.write_bytes(b"original")
    workbook = crabxl.Workbook()
    # All sheets removed is invalid for the writer. The target must survive.
    workbook.remove(workbook.active)
    with pytest.raises(IndexError):
        workbook.save(target)
    assert target.read_bytes() == b"original"
    assert list(tmp_path.iterdir()) == [target]


def test_sheet_names_and_order_match(engine):
    workbook = engine.Workbook()
    workbook.create_sheet("Sheet2")
    next_sheet = workbook.create_sheet()
    assert next_sheet.title == "Sheet3"
    next_sheet.title = "Data"
    workbook.create_sheet("data", index=0)
    assert workbook.sheetnames == ["data1", "Sheet", "Sheet2", "Data"]
    assert "Data" in workbook
    assert [sheet.title for sheet in workbook] == workbook.sheetnames
    workbook.remove(workbook["Sheet2"])
    assert workbook.sheetnames == ["data1", "Sheet", "Data"]


@pytest.mark.parametrize(
    "source,expected",
    [
        ("=A1+$B2+C$3+$D$4", "=B2+$B3+D$3+$D$4"),
        ("='A1'!A1+A1!B2", "='A1'!B2+A1!C3"),
        ("=T1[A1]+T2[[#Headers],[B2]]+A1", "=T1[A1]+T2[[#Headers],[B2]]+B2"),
        ('="A1"&"say ""B2"""+A1', '="A1"&"say ""B2"""+B2'),
        ("=LOG10(A1)+SUM(A1:B2:C3)", "=LOG10(B2)+SUM(B2:C3:D4)"),
        ("=XFD1048576", "=XFE1048577"),
        ("=AA1001001001+R1C1+named1", "=AA1001001001+R1C1+named1"),
    ],
)
def test_formula_translation_context(engine, source, expected):
    from importlib import import_module

    translator = import_module(engine.__name__ + ".formula.translate").Translator
    assert translator(source, "A1").translate_formula("B2") == expected


def test_formula_move_translation_and_failure_atomicity(engine):
    workbook = engine.Workbook()
    sheet = workbook.active
    sheet["B2"] = "=C3+$D$4"
    sheet.move_range("B2", rows=1, cols=1, translate=True)
    assert sheet["C3"].value == "=D4+$D$4"
    assert sheet["B2"].value is None
    if engine is crabxl:
        sheet["C4"] = "=A1"
        before = tuple(sheet.values)
        with pytest.raises(ValueError):
            sheet.move_range("C3:C4", rows=-1, translate=True)
        assert tuple(sheet.values) == before


def test_active_sheet_creation_and_loaded_catalog_match(engine, tmp_path):
    workbook = engine.Workbook()
    first = workbook.active
    second = workbook.create_sheet("Other")
    first["A1"] = 1
    second["A1"] = 2
    workbook.active = second
    assert workbook.active is second
    path = tmp_path / "active.xlsx"
    workbook.save(path)
    verified = openpyxl.load_workbook(path)
    assert verified.active.title == "Other"
    verified.close()
    loaded = engine.load_workbook(path)
    assert loaded.active.title == "Other" and loaded.active["A1"].value == 2
    loaded.close()


def test_owned_workbook_copy_order_removed_aliases_and_independent_values(
    engine, tmp_path
):
    workbook = engine.Workbook()
    source = workbook.active
    source.title = "Data"
    source.append([1, "text", "=A1+1"])
    source.append([])
    copied = workbook.copy_worksheet(source)
    assert copied.title == "Data Copy"
    assert copied.parent is workbook and copied["C1"].value == "=A1+1"
    copied["A1"] = 7
    assert source["A1"].value == 1
    assert workbook.copy_worksheet(source).title == "Data Copy1"
    workbook.move_sheet(copied, offset=-1)
    assert workbook.sheetnames == ["Data Copy", "Data", "Data Copy1"]
    assert workbook.index(source) == 1
    workbook.active = source
    cell = copied["A1"]
    workbook.remove(copied)
    assert cell.value == 7 and copied["A1"] is cell
    cell.value = 8
    assert copied["A1"].value == 8
    assert workbook.sheetnames == ["Data", "Data Copy1"]
    del workbook["Data Copy1"]
    assert workbook.active is None
    workbook.active = source
    target = tmp_path / "bank.xlsx"
    workbook.save(target)
    verified = openpyxl.load_workbook(target)
    assert verified.sheetnames == ["Data"] and verified.active["A1"].value == 1
    verified.close()
    with pytest.raises(ValueError):
        workbook.copy_worksheet(engine.Workbook().active)


@pytest.mark.parametrize("offset", [-10, -4, -2, -1, 0, 1, 2, 10])
def test_move_sheet_offsets_match_list_insertion(engine, offset):
    workbook = engine.Workbook()
    workbook.create_sheet("B")
    source = workbook.create_sheet("C")
    expected = workbook.worksheets
    old = expected.index(source)
    expected.remove(source)
    expected.insert(old + offset, source)
    workbook.move_sheet("C", offset=offset)
    assert workbook.worksheets == expected


def test_python_bank_aggregate_limit_atomic_copy_and_freed_space():
    workbook = crabxl.Workbook(max_memory_bytes=1600)
    first = workbook.active
    first["A1"] = 1
    second = workbook.create_sheet("B")
    second["A1"] = 2
    assert workbook._book.charged_bytes() == 1030
    with pytest.raises(MemoryError):
        workbook.copy_worksheet(first)
    assert workbook.sheetnames == ["Sheet", "B"]
    first["A2"] = 3
    second["A2"] = 4
    before = workbook._book.charged_bytes()
    with pytest.raises(MemoryError):
        second["A3"] = 5
    assert workbook._book.charged_bytes() == before and not second._native.contains(
        2, 0
    )
    with pytest.raises(MemoryError):
        first.insert_rows(1)
    assert first["A1"].value == 1
    workbook.remove(first)
    assert first["A1"].value == 1
    # Detached models remain caller-owned, outside the bank's allowance.
    second["A3"] = 5
    assert second["A3"].value == 5
    first["A1"] = 9
    assert first["A1"].value == 9


def test_native_bank_concurrent_copies_keep_owned_handles_valid():
    from concurrent.futures import ThreadPoolExecutor

    from crabxl._native import NativeBook

    book = NativeBook(10_000_000)
    source = book.create_sheet("Source")
    for row in range(400):
        source.append([("int", str(row))])

    def copies(worker):
        for iteration in range(10):
            copied = book.copy_sheet(source, f"Copy{worker}-{iteration}")
            assert copied.get(0, 0) == ("n", 0)
            assert copied.get(399, 0) == ("n", 399)
            book.remove_sheet(copied)
            assert copied.get(399, 0) == ("n", 399)
        return worker

    with ThreadPoolExecutor(max_workers=4) as pool:
        assert list(pool.map(copies, range(4))) == [0, 1, 2, 3]
    assert source.get(399, 0) == ("n", 399)
    assert book.charged_bytes() < 10_000_000


def test_loaded_calculation_chain_edit_removes_derived_parts(engine, tmp_path):
    source = tmp_path / "chain.xlsx"
    book = openpyxl.Workbook()
    book.active.append([1, "=A1+1"])
    book.save(source)
    with zipfile.ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        b"</Types>",
        b'<Override PartName="/custom/order.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/></Types>',
    )
    parts["xl/_rels/workbook.xml.rels"] = parts["xl/_rels/workbook.xml.rels"].replace(
        b"</Relationships>",
        b'<Relationship Id="chain" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain" Target="../custom/order.xml"/></Relationships>',
    )
    parts["custom/order.xml"] = (
        b'<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><c r="B1" i="1"/></calcChain>'
    )
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            archive.writestr(name, value)
    loaded = engine.load_workbook(source)
    loaded.active["A1"] = 7
    for filename in ["chain-edited.xlsx", "chain-repeat.xlsx"]:
        target = tmp_path / filename
        loaded.save(target)
        with zipfile.ZipFile(target) as archive:
            assert "custom/order.xml" not in archive.namelist()
            assert b"calcChain" not in archive.read("[Content_Types].xml")
            assert b"calcChain" not in archive.read("xl/_rels/workbook.xml.rels")
        verified = openpyxl.load_workbook(target)
        assert (
            verified.active["A1"].value == 7 and verified.active["B1"].value == "=A1+1"
        )
        verified.close()
    loaded.close()


@pytest.mark.parametrize(
    "value",
    [
        "_x0041_",
        "_x005F_x0041_",
        "_x005F__x0041_",
        "_x000D_",
        "_xD83D__xDE00_",
        "😀_x005f_<&>\r\n ",
    ],
)
def test_inline_ooxml_looking_literals_preserve_reference_spelling(
    engine, value, tmp_path
):
    book = engine.Workbook()
    book.active["A1"] = value
    assert book.active["A1"].value == value
    target = tmp_path / "literal.xlsx"
    book.save(target)
    checked = openpyxl.load_workbook(target)
    assert checked.active["A1"].value == value
    checked.close()
    own = engine.load_workbook(target)
    assert own.active["A1"].value == value
    own.active["B2"] = value
    own.save(tmp_path / "edited.xlsx")
    own.close()
    checked = openpyxl.load_workbook(tmp_path / "edited.xlsx")
    assert checked.active["A1"].value == checked.active["B2"].value == value
    checked.close()


@pytest.mark.parametrize("value", ["literal", "  a & b  ", "_x005F_x0041_", "", "🦀"])
def test_plain_shared_string_read_edit_and_repeat_save(engine, value, tmp_path):
    from xml.sax.saxutils import escape

    source = tmp_path / "shared-source.xlsx"
    initial = openpyxl.Workbook()
    initial.active["A1"] = "placeholder"
    initial.active["B1"] = "placeholder"
    initial.save(source)
    with zipfile.ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    parts["xl/sharedStrings.xml"] = (
        f'<sst xmlns="{main}" uniqueCount="1"><si><t>{escape(value)}</t></si></sst>'.encode()
    )
    parts["xl/_rels/workbook.xml.rels"] = parts["xl/_rels/workbook.xml.rels"].replace(
        b"</Relationships>",
        f'<Relationship Id="shared" Type="{rel}/sharedStrings" Target="sharedStrings.xml"/></Relationships>'.encode(),
    )
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        b"</Types>",
        b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>',
    )
    parts["xl/worksheets/sheet1.xml"] = (
        f'<worksheet xmlns="{main}"><dimension ref="A1:B1"/><sheetData><row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>0</v></c></row></sheetData></worksheet>'.encode()
    )
    with zipfile.ZipFile(source, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    expected = value.replace("x005F_", "")
    book = engine.load_workbook(source)
    assert list(book.active.values) == [(expected, expected)]
    book.active["A1"] = "changed"
    for index in range(2):
        target = tmp_path / f"shared-edited-{index}.xlsx"
        book.save(target)
        verified = openpyxl.load_workbook(target)
        assert verified.active["A1"].value == "changed"
        # The reference rewrites empty SST text as an absent inline literal.
        # Preserving core saves retain untouched source values instead.
        if value == "" and engine is openpyxl:
            assert verified.active["B1"].value is None
        else:
            assert verified.active["B1"].value == expected
        verified.close()
    book.close()


@pytest.mark.parametrize("location", ["inline", "shared"])
@pytest.mark.parametrize(
    "content,inline,shared",
    [
        (
            '<r><rPr><b/></rPr><t xml:space="preserve">  rich &amp; text </t></r><r><t>tail</t></r>',
            "  rich & text tail",
            "  rich & text tail",
        ),
        ("<r><rPr><b/></rPr><t>_x005F_x0041_</t></r>", "_x005F_x0041_", "_x0041_"),
        (
            "<r><rPr><b/></rPr><t>_x005F</t></r><r><t>_x0041_</t></r>",
            "_x005F_x0041_",
            "_x0041_",
        ),
        (
            '<r><t xml:space="preserve">  </t></r><r><rPr><i/></rPr><t></t></r>',
            "  ",
            "  ",
        ),
    ],
)
def test_default_rich_projection_edit_and_repeated_save(
    engine, location, content, inline, shared, tmp_path
):
    source = tmp_path / "rich-source.xlsx"
    initial = openpyxl.Workbook()
    initial.active["A1"] = initial.active["B1"] = "placeholder"
    initial.save(source)
    with zipfile.ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    if location == "shared":
        parts["xl/sharedStrings.xml"] = (
            f'<sst xmlns="{main}"><si>{content}</si></sst>'.encode()
        )
        parts["xl/_rels/workbook.xml.rels"] = parts[
            "xl/_rels/workbook.xml.rels"
        ].replace(
            b"</Relationships>",
            f'<Relationship Id="shared" Type="{rel}/sharedStrings" Target="sharedStrings.xml"/></Relationships>'.encode(),
        )
        parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
            b"</Types>",
            b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>',
        )

        def cell(address):
            return f'<c r="{address}" t="s"><v>0</v></c>'
    else:

        def cell(address):
            return f'<c r="{address}" t="inlineStr"><is>{content}</is></c>'

    parts["xl/worksheets/sheet1.xml"] = (
        f'<worksheet xmlns="{main}"><dimension ref="A1:B1"/><sheetData><row r="1">{cell("A1")}{cell("B1")}</row></sheetData></worksheet>'.encode()
    )
    with zipfile.ZipFile(source, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    expected = inline if location == "inline" else shared
    book = engine.load_workbook(source, rich_text=False)
    assert list(book.active.values) == [(expected, expected)]
    book.active["A1"] = "changed"
    for index in range(2):
        target = tmp_path / f"rich-edited-{index}.xlsx"
        book.save(target)
        checked = openpyxl.load_workbook(target, rich_text=False)
        assert list(checked.active.values) == [("changed", expected)]
        checked.close()
    book.close()


def test_preserving_rich_runs_is_separate_from_default_text_projection(tmp_path):
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    original = CellRichText(
        TextBlock(InlineFont(b=True, color="80445566"), "formatted"), " tail"
    )
    source = tmp_path / "preserving-rich.xlsx"
    reference = openpyxl.Workbook()
    reference.active["A1"] = original
    reference.active["B1"] = 1
    reference.save(source)
    book = crabxl.load_workbook(source)
    assert book.active["A1"].value == "formatted tail"
    book.active["B1"] = 2
    for index in range(2):
        target = tmp_path / f"preserved-{index}.xlsx"
        book.save(target)
        checked = openpyxl.load_workbook(target, rich_text=True)
        assert checked.active["A1"].value == original
        assert checked.active["B1"].value == 2
        checked.close()
    book.close()
    with pytest.raises(NotImplementedError):
        crabxl.load_workbook(source, rich_text=True)


@pytest.mark.parametrize("mac", [False, True], ids=["windows", "mac"])
@pytest.mark.parametrize("data_only", [False, True], ids=["formula", "cached"])
def test_loaded_numeric_dates_time_duration_and_formula_cache(
    engine, tmp_path, mac, data_only
):
    from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900

    source = tmp_path / "styled-numeric.xlsx"
    reference = openpyxl.Workbook()
    reference.epoch = CALENDAR_MAC_1904 if mac else CALENDAR_WINDOWS_1900
    sheet = reference.active
    for row, value in enumerate(
        [0, 0.5, 59, 60, 61, -0.5, 2958466, 45292.123456789], 1
    ):
        sheet.cell(row, 1, value).number_format = "yyyy-mm-dd hh:mm:ss.000"
    sheet["B1"] = 1.25
    sheet["B1"].number_format = "[h]:mm:ss.000"
    sheet["B2"] = True
    sheet["B2"].number_format = "yyyy-mm-dd"
    sheet["B3"] = "text"
    sheet["B3"].number_format = "yyyy-mm-dd"
    sheet["C1"] = "=1"
    sheet["C1"].number_format = "yyyy-mm-dd"
    reference.save(source)
    reference.close()
    with zipfile.ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(
        b"<f>1</f><v></v>", b"<f>1</f><v>61</v>"
    )
    assert b"<f>1</f><v>61</v>" in parts["xl/worksheets/sheet1.xml"]
    with zipfile.ZipFile(source, "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    with pytest.warns(UserWarning) if engine is openpyxl else nullcontext():
        book = engine.load_workbook(source, data_only=data_only)
    sheet = book.active
    assert sheet["A1"].value == time(0)
    assert sheet["A2"].value == time(12)
    assert sheet["A3"].value == (
        datetime(1904, 2, 29) if mac else datetime(1900, 2, 28)
    )
    assert sheet["A4"].value == (datetime(1904, 3, 1) if mac else datetime(1900, 2, 28))
    assert sheet["A5"].value == (datetime(1904, 3, 2) if mac else datetime(1900, 3, 1))
    assert sheet["A6"].value == (
        datetime(1903, 12, 31, 12) if mac else datetime(1899, 12, 29, 12)
    )
    assert sheet["A7"].value == "#VALUE!" and sheet["A7"].data_type == "e"
    assert sheet["A8"].value == (
        datetime(2028, 1, 2, 2, 57, 46, 667000)
        if mac
        else datetime(2024, 1, 1, 2, 57, 46, 667000)
    )
    assert sheet["B1"].value == timedelta(days=1, hours=6)
    assert sheet["B2"].value is True and sheet["B3"].value == "text"
    assert sheet["C1"].value == (sheet["A5"].value if data_only else "=1")
    book.close()


@pytest.mark.parametrize("reference", [openpyxl, crabxl], ids=["openpyxl", "crabxl"])
@pytest.mark.parametrize(
    "format_code,expected",
    [
        ("hh:mm:ss.000", time(2, 57, 46, 667000)),
        ("[h]:mm:ss.000", timedelta(hours=2, minutes=57, seconds=46, milliseconds=667)),
    ],
)
def test_loaded_fractional_clock_duration_baseline_rounding(
    reference, tmp_path, format_code, expected
):
    source = tmp_path / "fractional-clock.xlsx"
    book = openpyxl.Workbook()
    book.active["A1"] = 0.123456789
    book.active["A1"].number_format = format_code
    book.save(source)
    book.close()
    loaded = reference.load_workbook(source)
    try:
        assert loaded.active["A1"].value == expected
    finally:
        loaded.close()


@pytest.mark.parametrize(
    "value",
    [
        datetime(2024, 2, 29, 12, 3, 4, 123456),
        datetime(1899, 12, 31, 12, 0, 0, 123456),
        time(2, 57, 46, 666570),
        timedelta(microseconds=-1),
        timedelta(days=999999999, seconds=86399, microseconds=999999),
    ],
)
def test_literal_date_clock_duration_precision_matches_reference(
    engine, tmp_path, value
):
    book = engine.Workbook()
    book.active["A1"] = value
    assert book.active["A1"].value == value
    assert book.active["A1"].data_type == "d"
    path = tmp_path / "literal-precision.xlsx"
    book.save(path)
    reference = openpyxl.Workbook()
    reference.active["A1"] = value
    expected_path = tmp_path / "expected-precision.xlsx"
    reference.save(expected_path)
    reference.close()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        expected = openpyxl.load_workbook(expected_path)
        loaded = openpyxl.load_workbook(path)
    assert loaded.active["A1"].value == expected.active["A1"].value
    loaded.close()
    expected.close()
    book.close()


@pytest.mark.parametrize("iso", [False, True], ids=["numeric", "iso"])
@pytest.mark.parametrize("mac", [False, True], ids=["windows", "mac"])
def test_date_only_iso_creation_epoch_and_loaded_types(engine, tmp_path, iso, mac):
    from datetime import date

    from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900

    book = engine.Workbook(iso_dates=iso)
    assert book.iso_dates == iso
    assert book.epoch == CALENDAR_WINDOWS_1900
    book.epoch = CALENDAR_MAC_1904 if mac else CALENDAR_WINDOWS_1900
    values = (
        date(2024, 2, 29),
        datetime(2024, 2, 29, 12, 3, 4, 123456),
        time(12, 3, 4, 123456),
        timedelta(days=1, hours=6),
    )
    book.active.append(values)
    assert list(book.active.values) == [values]
    assert [book.active.cell(1, column).data_type for column in range(1, 5)] == [
        "d"
    ] * 4
    with pytest.raises(ValueError):
        book.epoch = 1904
    path = tmp_path / "date-mode.xlsx"
    book.save(path)
    loaded = engine.load_workbook(path)
    assert loaded.epoch == book.epoch
    expected = (
        date(2024, 2, 29) if iso else datetime(2024, 2, 29),
        datetime(2024, 2, 29, 12, 3, 4, 123000),
        time(12, 3, 4, 123000),
        timedelta(days=1, hours=6),
    )
    assert list(loaded.active.values) == [expected]
    assert type(loaded.active["A1"].value) is (date if iso else datetime)
    loaded.close()
    book.close()


def test_loaded_date_policy_changes_are_explicit_until_bank_integration(tmp_path):
    path = tmp_path / "loaded-policy.xlsx"
    book = openpyxl.Workbook()
    book.active["A1"] = datetime(2024, 1, 1)
    book.save(path)
    book.close()
    loaded = crabxl.load_workbook(path)
    try:
        with pytest.raises(NotImplementedError):
            loaded.iso_dates = True
        with pytest.raises(NotImplementedError):
            loaded.epoch = datetime(1904, 1, 1)
        assert loaded.iso_dates is False and loaded.epoch == datetime(1899, 12, 30)
    finally:
        loaded.close()


@pytest.mark.parametrize(
    "value, format_code",
    [
        (date(2024, 1, 2), "yyyy-mm-dd"),
        (datetime(2024, 1, 2, 3, 4, 5, 678900), "yyyy-mm-dd h:mm:ss"),
        (time(3, 4, 5, 678900), "h:mm:ss"),
        (timedelta(days=2, seconds=3, microseconds=678900), "[hh]:mm:ss"),
    ],
)
def test_saved_temporal_default_formats_match_public_assignments(
    engine, value, format_code, tmp_path
):
    book = engine.Workbook()
    book.active["A1"] = value
    assert book.active["A1"].value == value
    path = tmp_path / "temporal-format.xlsx"
    book.save(path)
    book.close()
    loaded = openpyxl.load_workbook(path)
    assert loaded.active["A1"].number_format == format_code
    loaded.close()


def test_replacing_temporal_value_retains_format_across_repeated_save(engine, tmp_path):
    initial_values = [
        (date(2024, 1, 2), "yyyy-mm-dd"),
        (datetime(2024, 1, 2, 3, 4, 5), "yyyy-mm-dd h:mm:ss"),
        (time(3, 4, 5), "h:mm:ss"),
        (timedelta(days=2, seconds=3), "[hh]:mm:ss"),
    ]
    replacements = [
        date(2024, 2, 3),
        datetime(2024, 2, 3, 4, 5, 6),
        time(4, 5, 6),
        timedelta(days=3, seconds=4),
        42,
    ]
    cases = [
        (initial, code, replacement)
        for initial, code in initial_values
        for replacement in replacements
    ]
    book = engine.Workbook()
    reference = openpyxl.Workbook()
    for row, (initial, _, _) in enumerate(cases, 1):
        book.active.cell(row, 1, initial)
        reference.active.cell(row, 1, initial)
    first = tmp_path / "first.xlsx"
    book.save(first)
    for row, (_, _, replacement) in enumerate(cases, 1):
        book.active.cell(row, 1).value = replacement
        assert book.active.cell(row, 1).value == replacement, cases[row - 1]
        reference.active.cell(row, 1).value = replacement
    expected_path = tmp_path / "expected.xlsx"
    reference.save(expected_path)
    reference.close()
    expected = openpyxl.load_workbook(expected_path)
    expected_values = []
    for row, (_, code, _) in enumerate(cases, 1):
        assert expected.active.cell(row, 1).number_format == code, cases[row - 1]
        expected_values.append(expected.active.cell(row, 1).value)
    expected.close()
    for index in range(2):
        path = tmp_path / f"replaced-{index}.xlsx"
        book.save(path)
        loaded = openpyxl.load_workbook(path)
        for row, (_, code, _) in enumerate(cases, 1):
            cell = loaded.active.cell(row, 1)
            assert cell.number_format == code, cases[row - 1]
            assert cell.value == expected_values[row - 1], cases[row - 1]
        loaded.close()
    original = openpyxl.load_workbook(first)
    for row, (_, code, _) in enumerate(cases, 1):
        assert original.active.cell(row, 1).number_format == code, cases[row - 1]
    original.close()
    book.close()
