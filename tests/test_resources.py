"""Resource controls must affect native behavior in both public read modes."""

import os
from pathlib import Path

import crabxl
import openpyxl
import pytest

MIB = 1024 * 1024


def test_resource_validation_modes_and_auto_controls(tmp_path):
    for factory in [crabxl.AutoMemory, crabxl.ResourceLimits]:
        argument = (
            "available_bytes" if factory is crabxl.AutoMemory else "max_cell_bytes"
        )
        for value, error in [
            (True, TypeError),
            (1.5, TypeError),
            (-1, ValueError),
            (2**64, ValueError),
        ]:
            with pytest.raises(error):
                factory(**{argument: value})
    for argument, value in [
        ("fraction_per_mille", 0),
        ("fraction_per_mille", 1001),
        ("concurrent_operations", 0),
        ("concurrent_operations", 65536),
    ]:
        with pytest.raises(ValueError):
            crabxl.AutoMemory(**{argument: value})
    with pytest.raises(TypeError):
        crabxl.ResourceOptions(limits={"max_cell_bytes": 1})
    with pytest.raises(TypeError):
        crabxl.SharedStringOptions(cache_bytes=True)
    with pytest.raises(ValueError):
        crabxl.SharedStringOptions(storage="unknown")
    with pytest.raises(ValueError, match="Disk/cache"):
        crabxl.SharedStringOptions(storage="memory", temp_directory=tmp_path)
    path = tmp_path / "absent.xlsx"
    for read_only, options in [
        (False, crabxl.ResourceOptions(limits=crabxl.ResourceLimits(max_batch_rows=1))),
        (
            True,
            crabxl.ResourceOptions(
                limits=crabxl.ResourceLimits(max_materialized_bytes=1)
            ),
        ),
        (True, crabxl.ResourceOptions(max_patch_cells=1)),
    ]:
        with pytest.raises(ValueError):
            crabxl.load_workbook(path, read_only=read_only, resource_options=options)
    with pytest.raises(ValueError, match="Choose"):
        crabxl.load_workbook(
            path,
            max_memory_bytes=MIB,
            resource_options=crabxl.ResourceOptions(auto_memory=crabxl.AutoMemory()),
        )
    with pytest.raises(ValueError, match="Choose"):
        crabxl.Workbook(max_memory_bytes=MIB, auto_memory=crabxl.AutoMemory())
    with pytest.raises(MemoryError):
        crabxl.Workbook(auto_memory=crabxl.AutoMemory(maximum_bytes=1))
    for write_only in [False, True]:
        book = crabxl.Workbook(
            write_only=write_only,
            auto_memory=crabxl.AutoMemory(
                available_bytes=64 * MIB,
                headroom_bytes=0,
                maximum_bytes=8 * MIB,
                concurrent_operations=2,
            ),
        )
        sheet = book.create_sheet() if write_only else book.active
        sheet.append([7, "🦀", True])
        output = tmp_path / f"auto-{write_only}.xlsx"
        book.save(output, compression_level=1)
        book.close()
        checked = openpyxl.load_workbook(output, read_only=True)
        assert list(checked.active.values) == [(7, "🦀", True)]
        checked.close()


def test_read_edit_and_batch_limits_are_enforced_without_corrupting_source(tmp_path):
    source = tmp_path / "source.xlsx"
    original = openpyxl.Workbook()
    original.active.append([1, "long text", 3])
    original.save(source)
    original.close()
    before = source.read_bytes()
    for read_only in [False, True]:
        for limits in [
            crabxl.ResourceLimits(max_cell_bytes=3),
            crabxl.ResourceLimits(max_row_cells=1),
        ]:
            book = crabxl.load_workbook(
                source,
                read_only=read_only,
                max_memory_bytes=8 * MIB,
                resource_options=crabxl.ResourceOptions(limits=limits),
            )
            try:
                with pytest.raises(ValueError):
                    list(book.active.values)
            finally:
                book.close()
        with pytest.raises(ValueError):
            crabxl.load_workbook(
                source,
                read_only=read_only,
                resource_options=crabxl.ResourceOptions(
                    limits=crabxl.ResourceLimits(max_archive_entries=1)
                ),
            )
    book = crabxl.load_workbook(
        source,
        max_memory_bytes=8 * MIB,
        resource_options=crabxl.ResourceOptions(
            limits=crabxl.ResourceLimits(max_materialized_bytes=1)
        ),
    )
    try:
        with pytest.raises(MemoryError):
            list(book.active.values)
    finally:
        book.close()
    book = crabxl.load_workbook(
        source,
        max_memory_bytes=8 * MIB,
        resource_options=crabxl.ResourceOptions(max_patch_cells=1),
    )
    try:
        book.active["A1"] = 99
        with pytest.raises(MemoryError):
            book.active["B1"] = 27
        assert book.active["A1"].value == 99
        assert book.active["B1"].value == "long text"
        book.active["A1"] = 88
        assert book.active["A1"].value == 88
    finally:
        book.close()
    book = crabxl.load_workbook(
        source,
        read_only=True,
        max_memory_bytes=8 * MIB,
        resource_options=crabxl.ResourceOptions(
            limits=crabxl.ResourceLimits(max_batch_rows=1, max_batch_bytes=4096)
        ),
    )
    assert list(book.active.values) == [(1, "long text", 3)]
    book.close()
    assert source.read_bytes() == before


def test_shared_string_placement_thresholds_limits_and_cleanup(
    tmp_path, shared_strings_source
):
    values = [f"v-{index:04}-" + "x" * 96 for index in range(10000)]
    source = shared_strings_source(tmp_path / "strings.xlsx", values)
    temporary = tmp_path / "owned-sst"
    temporary.mkdir()
    for read_only in [False, True]:
        for storage in ["memory", "auto", "disk"]:
            strings = crabxl.SharedStringOptions(
                storage=storage,
                memory_bytes=(8 if storage == "memory" else 2) * MIB,
                **(
                    {}
                    if storage == "memory"
                    else {"cache_bytes": 4096, "temp_directory": temporary}
                ),
            )
            book = crabxl.load_workbook(
                source,
                read_only=read_only,
                max_memory_bytes=4 * MIB,
                resource_options=crabxl.ResourceOptions(shared_strings=strings),
            )
            iterator = book.active.values
            assert next(iterator) == (values[0],)
            descriptors = Path("/proc/self/fd")
            if descriptors.is_dir():
                owned = []
                for descriptor in descriptors.iterdir():
                    try:
                        if str(temporary) in os.readlink(descriptor):
                            owned.append(descriptor)
                    except OSError:
                        pass
                assert len(owned) == (0 if storage == "memory" else 2)
            assert list(iterator) == [(value,) for value in values[1:]]
            book.close()
            assert not list(temporary.iterdir())
        for strings, error in [
            (
                crabxl.SharedStringOptions(storage="memory", memory_bytes=2 * MIB),
                MemoryError,
            ),
            (
                crabxl.SharedStringOptions(
                    storage="disk", max_temp_bytes=0, temp_directory=temporary
                ),
                ValueError,
            ),
            (crabxl.SharedStringOptions(max_entries=1), ValueError),
            (
                crabxl.SharedStringOptions(
                    storage="disk", temp_directory=temporary / "absent"
                ),
                OSError,
            ),
        ]:
            book = crabxl.load_workbook(
                source,
                read_only=read_only,
                max_memory_bytes=4 * MIB,
                resource_options=crabxl.ResourceOptions(shared_strings=strings),
            )
            try:
                with pytest.raises(error):
                    list(book.active.values)
            finally:
                book.close()
            assert not list(temporary.iterdir())
