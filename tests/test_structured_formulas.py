"""Run structured formula calls against the pinned public reference and adapter."""
from importlib import import_module
import gc
import weakref
import zipfile
import openpyxl
import crabxl
import pytest


@pytest.fixture(params=[openpyxl, crabxl], ids=["openpyxl", "crabxl"])
def engine(request):
    return request.param


@pytest.mark.parametrize("kind", ["array", "table"])
def test_structured_creation_mutation_and_readback(engine, kind, tmp_path):
    module = import_module(engine.__name__ + ".worksheet.formula")
    book = engine.Workbook()
    sheet = book.active
    value = module.ArrayFormula("A1:B2", "=SUM(C1:C2)") if kind == "array" else module.DataTableFormula("A1:B2", dt2D=True, r1="C1", r2="D1")
    sheet["A1"] = value
    assert sheet["A1"].data_type == "f"
    assert vars(sheet["A1"].value) == vars(value)
    attached = sheet["A1"].value
    attached.ref = "$A$1:$B$2"
    if kind == "array":
        attached.text = "=SUM(D1:D2)"
    else:
        attached.r1 = "$D$1"
        attached.del1 = True
    assert vars(sheet["A1"].value) == vars(attached)
    first = tmp_path / "first.xlsx"
    book.save(first)
    book.close()
    loaded = openpyxl.load_workbook(first)
    value = loaded.active["A1"].value
    assert value.ref == "$A$1:$B$2"
    if kind == "array":
        assert value.text == "=SUM(D1:D2)"
    else:
        assert value.r1 == "$D$1"
        assert value.del1 == "1"
    loaded.close()
    loaded = engine.load_workbook(first)
    value = loaded.active["A1"].value
    assert value.ref == "$A$1:$B$2"
    value.ref = "A1:B3"
    second = tmp_path / "second.xlsx"
    loaded.save(second)
    loaded.close()
    check = openpyxl.load_workbook(second)
    assert check.active["A1"].value.ref == "A1:B3"
    check.close()


def test_formula_prefix_is_removed_exactly_once(engine, tmp_path):
    book = engine.Workbook()
    book.active["A1"] = "==1"
    assert book.active["A1"].value == "==1"
    path = tmp_path / "prefix.xlsx"
    book.save(path)
    book.close()
    loaded = engine.load_workbook(path)
    assert loaded.active["A1"].value == "==1"
    loaded.close()


@pytest.mark.parametrize("cached", [False, True])
def test_shared_and_empty_caches_match_public_reference(engine, cached, tmp_path):
    path = tmp_path / "source.xlsx"
    source = openpyxl.Workbook()
    source.save(path)
    source.close()
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/worksheets/sheet1.xml"] = b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="B1"><f t="shared" si="4294967295" ref="A1:D2">B1+$B$2+C3</f><v>0</v></c><c r="D1"><f t="shared" si="4294967295"/><v>1</v></c><c r="E1" t="str"><f>1</f><v/></c></row></sheetData></worksheet>'
    with zipfile.ZipFile(path, "w") as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)
    book = engine.load_workbook(path, data_only=cached)
    assert [book.active[cell].value for cell in ("B1", "D1", "E1")] == ([0, 1, None] if cached else ["=B1+$B$2+C3", "=D1+$B$2+E3", "=1"])
    book.close()


def test_adapter_formula_views_do_not_keep_workbooks_alive():
    from crabxl.worksheet.formula import ArrayFormula
    book = crabxl.Workbook()
    sheet = book.active
    sheet["A1"] = ArrayFormula("A1:B2", "=1")
    value = sheet["A1"].value
    sheet_ref = weakref.ref(sheet)
    book.close()
    del sheet, book
    gc.collect()
    assert sheet_ref() is None
    value.text = "=2"
    assert value.text == "=2"


@pytest.mark.parametrize("flags", [
    {},
    {"dt2D": True, "dtr": False, "del1": False, "del2": True, "r1": "C1", "r2": "D1"},
    {"dt2D": False, "dtr": False, "r1": ""},
    {"dt2D": "0", "dtr": "false", "del1": "0"},
])
def test_all_data_table_properties_match_public_save_reload(engine, flags, tmp_path):
    module = import_module(engine.__name__ + ".worksheet.formula")
    reference = openpyxl.Workbook()
    reference.active["A1"] = openpyxl.worksheet.formula.DataTableFormula(ref="A1:B2", **flags)
    reference_path = tmp_path / "reference.xlsx"
    reference.save(reference_path)
    reference.close()
    check = openpyxl.load_workbook(reference_path)
    expected = vars(check.active["A1"].value)
    check.close()
    book = engine.Workbook()
    book.active["A1"] = module.DataTableFormula(ref="A1:B2", **flags)
    output = tmp_path / "native.xlsx"
    book.save(output)
    book.close()
    loaded = engine.load_workbook(output)
    assert vars(loaded.active["A1"].value) == expected
    loaded.active["A1"].value.ref = "A1:B3"
    edited = tmp_path / "edited.xlsx"
    loaded.save(edited)
    loaded.close()
    check = openpyxl.load_workbook(edited)
    assert vars(check.active["A1"].value) == {**expected, "ref": "A1:B3"}
    check.close()

@pytest.mark.parametrize("content, cached, expected", [
    ('<c r="A1"><f t="future" ref="invalid">1+1</f><v>2</v></c>', True, 2),
    ('<c r="A1" cm="1"><f t="array" ref="A1">_xlfn.SEQUENCE(1)</f><v>7</v></c>', True, 7),
    ('<c r="A1" vm="not-an-index"><v>9</v></c>', False, 9),
])
def test_visible_metadata_and_discarded_formula_semantics(engine, content, cached, expected, tmp_path):
    path = tmp_path / "projected.xlsx"
    original = openpyxl.Workbook()
    original.save(path)
    original.close()
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/worksheets/sheet1.xml"] = ('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1">' + content + '</row></sheetData></worksheet>').encode()
    with zipfile.ZipFile(path, "w") as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)
    book = engine.load_workbook(path, data_only=cached)
    assert book.active["A1"].value == expected
    book.close()

@pytest.mark.parametrize('reference', [None, 'A1:B2'])
@pytest.mark.parametrize('text', [None, '', '=', '=1', '1', 'abc', '==1', '\u03b1x'])
def test_literal_array_text_presence_and_source_body_match_public_calls(engine, reference, text, tmp_path):
    module = import_module(engine.__name__ + '.worksheet.formula')
    book = engine.Workbook()
    book.active['A1'] = module.ArrayFormula(reference, text)
    assert vars(book.active['A1'].value) == {'ref': reference, 'text': text}
    path = tmp_path / 'literal-array.xlsx'
    book.save(path)
    assert vars(book.active['A1'].value) == {'ref': reference, 'text': text}
    book.save(path)
    book.close()
    loaded = engine.load_workbook(path)
    assert vars(loaded.active['A1'].value) == {'ref': reference, 'text': '=' + (text or '')[1:]}
    loaded.close()
    cached = engine.load_workbook(path, data_only=True)
    assert cached.active['A1'].value is None
    cached.close()


@pytest.mark.parametrize('text', [None, '', 'abc', '\u03b1x'])
def test_loaded_array_text_property_updates_use_canonical_literal_conversion(engine, text, tmp_path):
    module = import_module(engine.__name__ + '.worksheet.formula')
    book = engine.Workbook()
    book.active['A1'] = module.ArrayFormula('A1:B2', '=1')
    path = tmp_path / 'loaded-array-text.xlsx'
    book.save(path)
    book.close()
    loaded = engine.load_workbook(path)
    value = loaded.active['A1'].value
    value.text = text
    assert loaded.active['A1'].value.text == text
    loaded.save(path)
    loaded.save(path)
    loaded.close()
    reloaded = engine.load_workbook(path)
    assert reloaded.active['A1'].value.text == '=' + (text or '')[1:]
    reloaded.close()


@pytest.mark.parametrize('reference', ['', '$A$1:$B$2', 'Sheet1!A1:B2', 'not a range'])
def test_literal_array_reference_properties_and_loaded_updates(engine, reference, tmp_path):
    module = import_module(engine.__name__ + '.worksheet.formula')
    book = engine.Workbook()
    book.active['A1'] = module.ArrayFormula(reference, '=1')
    assert book.active['A1'].value.ref == reference
    path = tmp_path / 'literal-reference.xlsx'
    book.save(path)
    assert book.active['A1'].value.ref == reference
    book.close()
    loaded = engine.load_workbook(path)
    assert loaded.active['A1'].value.ref == (reference or None)
    loaded.active['A1'].value.ref = 'Sheet2!C1:D2'
    loaded.save(path)
    loaded.close()
    checked = openpyxl.load_workbook(path)
    assert checked.active['A1'].value.ref == 'Sheet2!C1:D2'
    checked.close()


@pytest.mark.parametrize('reference', ['A1:B2', 'Sheet1!A1:B2', 'opaque'])
@pytest.mark.parametrize('input1', [None, '', 'C1', 'Sheet1!C1', 'input'])
def test_literal_table_references_and_inputs_match_public_save(engine, reference, input1, tmp_path):
    module = import_module(engine.__name__ + '.worksheet.formula')
    book = engine.Workbook()
    book.active['A1'] = module.DataTableFormula(reference, r1=input1, r2='')
    assert book.active['A1'].value.ref == reference
    assert book.active['A1'].value.r1 == input1
    path = tmp_path / 'literal-table.xlsx'
    book.save(path)
    book.close()
    loaded = engine.load_workbook(path)
    assert loaded.active['A1'].value.ref == reference
    assert loaded.active['A1'].value.r1 == (input1 or None)
    assert loaded.active['A1'].value.r2 is None
    loaded.active['A1'].value.r1 = 'updated input'
    loaded.save(path)
    loaded.close()
    checked = openpyxl.load_workbook(path)
    assert checked.active['A1'].value.r1 == 'updated input'
    checked.close()


@pytest.mark.parametrize('literal', ['', 'false', '0', 'invalid', ' true ', 'quoted & value'])
def test_literal_table_flags_preserve_public_properties_and_save_rules(engine, literal, tmp_path):
    module = import_module(engine.__name__ + '.worksheet.formula')
    fields = ['ca', 'dt2D', 'dtr', 'del1', 'del2']
    book = engine.Workbook()
    book.active['A1'] = module.DataTableFormula('A1:B2', **dict.fromkeys(fields, literal))
    assert all(getattr(book.active['A1'].value, name) == literal for name in fields)
    path = tmp_path / 'literal-flags.xlsx'
    book.save(path)
    book.close()
    loaded = engine.load_workbook(path)
    assert all(getattr(loaded.active['A1'].value, name) == (literal if literal else False) for name in fields)
    loaded.active['A1'].value.ca = 'updated flag'
    loaded.save(path)
    loaded.close()
    checked = openpyxl.load_workbook(path)
    assert checked.active['A1'].value.ca == 'updated flag'
    checked.close()


@pytest.mark.parametrize('attributes, expected', [
    ('t="future" ref="invalid"', '=A1+1'),
    ('ca="invalid" si="invalid" r1="opaque" t="normal"', '=A1+1'),
    ('unknown="value" t="normal"', '=A1+1'),
    ('ca="invalid" unknown="value" si="3" ref="opaque" t="shared"', '=A1+1'),
    ('t="array" ref="opaque" aca="invalid" r1="unused" unknown="value"', {'ref': 'opaque', 'text': '=A1+1'}),
    ('t="dataTable" ref="opaque" ca="invalid" unknown="value"', {'ref': 'opaque', 'ca': 'invalid', 'dt2D': False, 'dtr': False, 'r1': None, 'r2': None, 'del1': False, 'del2': False}),
])
@pytest.mark.parametrize('cached', [False, True])
def test_compatible_source_formula_headers_match_visible_public_values(engine, attributes, expected, cached, tmp_path):
    path = tmp_path / 'header.xlsx'
    source = openpyxl.Workbook()
    source.save(path)
    source.close()
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts['xl/worksheets/sheet1.xml'] = ('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1"><f ' + attributes + '>A1+1</f><v>2</v></c></row></sheetData></worksheet>').encode()
    with zipfile.ZipFile(path, 'w') as archive:
        for name, value in parts.items():
            archive.writestr(name, value)
    book = engine.load_workbook(path, data_only=cached)
    value = book.active['A1'].value
    assert (vars(value) if isinstance(expected, dict) and not cached else value) == (2 if cached else expected)
    book.close()


@pytest.mark.parametrize("master, follower", [
    (value, value) for value in (None, "", "0", "01", "1", "+1", "-1", "4294967296", "text", "1 ", "group & value")
] + [("01", "1"), ("1", "01"), (None, ""), ("", None)])
@pytest.mark.parametrize("cached", [False, True])
def test_literal_shared_identities_loading_and_unmodified_repeat_save(engine, master, follower, cached, tmp_path):
    from xml.sax.saxutils import quoteattr
    path = tmp_path / "shared-identity.xlsx"
    source = openpyxl.Workbook()
    source.save(path)
    source.close()
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    def index(value):
        return "" if value is None else " si=" + quoteattr(value)
    parts["xl/worksheets/sheet1.xml"] = (
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
        '<row r="1"><c r="A1"><f t="shared"' + index(master) + '>A1+1</f><v>2</v></c></row>'
        '<row r="2"><c r="A2"><f t="shared"' + index(follower) + '/><v>3</v></c></row>'
        '</sheetData></worksheet>'
    ).encode()
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            archive.writestr(name, value)
    formulas = ["=A1+1", "=A2+1" if master == follower else "="]
    expected = [2, 3] if cached else formulas
    book = engine.load_workbook(path, data_only=cached)
    assert [book.active.cell(row, 1).value for row in (1, 2)] == expected
    assert list(book.active.iter_rows(min_row=1, max_row=2, max_col=1, values_only=True)) == [(expected[0],), (expected[1],)]
    if not cached:
        for attempt in range(2):
            saved = tmp_path / f"shared-repeat-{attempt}.xlsx"
            book.save(saved)
            loaded = openpyxl.load_workbook(saved)
            assert [loaded.active.cell(row, 1).value for row in (1, 2)] == formulas
            loaded.close()
    book.close()
