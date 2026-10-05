"""Compare ordinary editable values iteration against the released adapter."""

import argparse
import hashlib
import json
import platform
import statistics
import subprocess
import tempfile
import time
from pathlib import Path


def identity(python):
    probe = """
import hashlib, importlib.metadata, json, platform
from pathlib import Path
import crabxl, crabxl._native
print(json.dumps({
    'python': platform.python_version(),
    'crabxl': importlib.metadata.version('crabxl'),
    'openpyxl': importlib.metadata.version('openpyxl'),
    'native_sha256': hashlib.sha256(Path(crabxl._native.__file__).read_bytes()).hexdigest(),
    'python_source_sha256': hashlib.sha256(Path(crabxl.__file__).read_bytes()).hexdigest(),
}))
"""
    return json.loads(subprocess.check_output([str(python), "-c", probe], text=True))


def worker(engine, path, read_only=False):
    module = __import__(engine)
    started = time.perf_counter()
    options = {"max_memory_bytes": 1024**3} if engine == "crabxl" else {}
    book = module.load_workbook(path, read_only=read_only, **options)
    opened = time.perf_counter()
    count = total = 0
    for row in book.active.iter_rows(values_only=True):
        count += len(row)
        total += sum(row)
    consumed = time.perf_counter()
    book.close()
    print(
        json.dumps(
            {
                "cells": count,
                "checksum": total,
                "open_seconds": opened - started,
                "iterate_seconds": consumed - opened,
            }
        )
    )


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--worker", choices=("crabxl", "openpyxl"))
    parser.add_argument("--path", type=Path)
    parser.add_argument("--before-python", type=Path)
    parser.add_argument("--after-python", type=Path)
    parser.add_argument("--measure", type=Path)
    parser.add_argument("--rows", type=int, nargs="+", default=[1000, 10000, 100000])
    parser.add_argument("--runs", type=int, default=5)
    parser.add_argument("--read-only", action="store_true")
    parser.add_argument("--skip-reference", action="store_true")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).parent / "results/values-calls.json",
    )
    args = parser.parse_args()
    if args.worker:
        worker(args.worker, args.path, args.read_only)
        return
    if not all((args.before_python, args.after_python, args.measure)):
        parser.error(
            "Specify both installed interpreters and the native measure helper"
        )
    import openpyxl

    report = {
        "semantics": ("Read-only" if args.read_only else "Ordinary editable")
        + " load_workbook and iter_rows(values_only=True). "
        "Same generated numeric cells, count and exact checksum. "
        "One warmup and rotating serial cold-process samples include imports and cleanup; "
        "generation/build excluded. Load/iteration boundaries are mode-specific; "
        "compare the complete operation, including deferred parsing or model materialization.",
        "platform": platform.platform(),
        "installed_packages": {
            "before": identity(args.before_python),
            "after": identity(args.after_python),
        },
        "reference_version": openpyxl.__version__,
        "model_allowance_bytes": 1024**3,
        "read_temporary_bytes": 0,
        "cases": [],
    }
    with tempfile.TemporaryDirectory(prefix="crabxl-values-benchmark-") as directory:
        for rows in args.rows:
            path = Path(directory) / f"numbers-{rows}.xlsx"
            book = openpyxl.Workbook(write_only=True)
            sheet = book.create_sheet("Sheet")
            for row in range(rows):
                sheet.append([row * 10 + column for column in range(10)])
            book.save(path)
            commands = {
                "before": (args.before_python, "crabxl"),
                "after": (args.after_python, "crabxl"),
                "openpyxl_normal": (args.after_python, "openpyxl"),
            }
            if args.skip_reference:
                del commands["openpyxl_normal"]
            elif args.read_only:
                commands["openpyxl_read_only"] = commands.pop("openpyxl_normal")
            samples = {name: [] for name in commands}
            names = list(commands)
            for index in range(args.runs + 1):
                for name in names[index % len(names) :] + names[: index % len(names)]:
                    python, engine = commands[name]
                    result = subprocess.run(
                        [
                            str(args.measure),
                            str(python),
                            str(Path(__file__).resolve()),
                            "--worker",
                            engine,
                            "--path",
                            str(path),
                            *(["--read-only"] if args.read_only else []),
                        ],
                        check=True,
                        capture_output=True,
                        text=True,
                    )
                    output = json.loads(result.stdout)
                    cells = rows * 10
                    assert output["cells"] == cells
                    assert output["checksum"] == cells * (cells - 1) // 2
                    sample = json.loads(result.stderr.split("MEASURE ")[-1])
                    sample.update(output)
                    if index:
                        samples[name].append(sample)
                    print(rows, name, sample["seconds"], flush=True)
            report["cases"].append(
                {
                    "rows": rows,
                    "columns": 10,
                    "file_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                    "samples": samples,
                    "medians": {
                        name: {
                            key: statistics.median(sample[key] for sample in values)
                            for key in (
                                "seconds",
                                "cpu_seconds",
                                "peak_rss_kib",
                                "open_seconds",
                                "iterate_seconds",
                            )
                        }
                        for name, values in samples.items()
                    },
                }
            )
            args.output.write_text(json.dumps(report, indent=2) + "\n")


if __name__ == "__main__":
    main()
