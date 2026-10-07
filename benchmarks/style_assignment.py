"""Complete public style edits and saves with fresh-process resource evidence."""

import argparse
import hashlib
import importlib
import json
import os
import resource
import statistics
import subprocess
import sys
import tempfile
import threading
import time
from datetime import date
from pathlib import Path


def peak_rss_kib():
    if sys.platform == "linux":
        for line in Path("/proc/self/status").read_text().splitlines():
            if line.startswith("VmHWM:"):
                return int(line.split()[1])
    peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    return peak // 1024 if sys.platform == "darwin" else peak


def worker(engine, mode, rows, source):
    with tempfile.TemporaryDirectory(prefix="crabxl-styles-") as directory:
        root = Path(directory)
        spools = root / "spools"
        spools.mkdir()
        os.environ["TMPDIR"] = str(spools)
        tempfile.tempdir = str(spools)
        module = importlib.import_module(engine)
        stop = threading.Event()
        peak_temp = [0]

        def sample():
            while not stop.wait(0.005):
                total = 0
                for path in spools.rglob("*"):
                    try:
                        if path.is_file():
                            total += path.stat().st_size
                    except FileNotFoundError:
                        pass
                peak_temp[0] = max(peak_temp[0], total)

        monitor = threading.Thread(target=sample, daemon=True)
        monitor.start()
        wall = time.perf_counter()
        cpu = time.process_time()
        if mode == "loaded":
            book = module.load_workbook(source)
            sheet = book.active
        else:
            book = module.Workbook(write_only=mode == "write_only")
            sheet = book.create_sheet() if mode == "write_only" else book.active
        if mode == "write_only":
            cell_type = importlib.import_module(engine + ".cell.cell").WriteOnlyCell
        for row in range(rows):
            values = [date(2024, 1, 2), *range(row * 10 + 1, row * 10 + 10)]
            if mode == "write_only":
                cells = [cell_type(sheet, value) for value in values]
                for column, cell in enumerate(cells):
                    cell.number_format = "General" if column == 0 else "0.0000"
                sheet.append(cells)
            else:
                if mode == "owned":
                    sheet.append(values)
                for column in range(10):
                    sheet.cell(row + 1, column + 1).number_format = (
                        "General" if column == 0 else "0.0000"
                    )
        output = root / "output.xlsx"
        book.save(output)
        book.close()
        seconds = time.perf_counter() - wall
        cpu_seconds = time.process_time() - cpu
        rss = peak_rss_kib()
        stop.set()
        monitor.join()
        remaining = sum(
            path.stat().st_size for path in spools.rglob("*") if path.is_file()
        )
        assert remaining == 0
        import openpyxl

        checked = openpyxl.load_workbook(output, read_only=True)
        checksum = count = 0
        for row, cells in enumerate(checked.active.iter_rows()):
            assert len(cells) == 10
            for column, cell in enumerate(cells):
                expected = 45293 if column == 0 else row * 10 + column
                assert cell.value == expected and type(cell.value) is int
                assert cell.number_format == ("General" if column == 0 else "0.0000")
                checksum += expected
                count += 1
        assert count == rows * 10
        checked.close()
        native = (
            Path(importlib.import_module("crabxl._native").__file__)
            if engine == "crabxl"
            else None
        )
        return {
            "engine": engine,
            "mode": mode,
            "rows": rows,
            "cells": count,
            "seconds": seconds,
            "cpu_seconds": cpu_seconds,
            "peak_rss_kib": rss,
            "sampled_peak_temp_bytes": peak_temp[0],
            "remaining_temp_bytes": remaining,
            "output_bytes": output.stat().st_size,
            "checksum": checksum,
            "native_sha256": hashlib.sha256(native.read_bytes()).hexdigest()
            if native
            else None,
        }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--worker", choices=("crabxl", "openpyxl"))
    parser.add_argument(
        "--mode", choices=("owned", "loaded", "write_only"), default="owned"
    )
    parser.add_argument("--rows", type=int, default=10000)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--source", type=Path)
    args = parser.parse_args()
    if args.worker:
        print(json.dumps(worker(args.worker, args.mode, args.rows, args.source)))
        return

    import openpyxl

    directory = tempfile.TemporaryDirectory(prefix="crabxl-style-source-")
    source = Path(directory.name) / "source.xlsx"
    seed = openpyxl.Workbook(write_only=True)
    sheet = seed.create_sheet()
    for row in range(args.rows):
        sheet.append([date(2024, 1, 2), *range(row * 10 + 1, row * 10 + 10)])
    seed.save(source)
    seed.close()

    def run(engine, mode):
        return json.loads(
            subprocess.check_output(
                [
                    sys.executable,
                    __file__,
                    "--worker",
                    engine,
                    "--mode",
                    mode,
                    "--rows",
                    str(args.rows),
                    "--source",
                    str(source),
                ],
                text=True,
            )
        )

    for mode in ("owned", "loaded", "write_only"):
        for engine in ("openpyxl", "crabxl"):
            run(engine, mode)
    samples = []
    for index in range(3):
        for mode in ("owned", "loaded", "write_only"):
            for engine in (
                ("openpyxl", "crabxl") if index % 2 == 0 else ("crabxl", "openpyxl")
            ):
                result = run(engine, mode)
                samples.append(result)
                print(json.dumps(result), flush=True)
    summary = {
        f"{engine}:{mode}": {
            field: statistics.median(
                sample[field]
                for sample in samples
                if sample["engine"] == engine and sample["mode"] == mode
            )
            for field in (
                "seconds",
                "cpu_seconds",
                "peak_rss_kib",
                "sampled_peak_temp_bytes",
            )
        }
        for engine in ("openpyxl", "crabxl")
        for mode in ("owned", "loaded", "write_only")
    }
    record = {
        "status": "Unreleased A9 binding preview, not the public A8 package",
        "core_revision": "84b7aa06e587ce1c4f7f5779d954bdef4a7788a4",
        "script_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        "measurement": "One warm-up and three alternating fresh-process samples. Owned/write-only include creation, conversion, all style assignments and ZIP finalization; loaded includes load, all style assignments and finalization. Feature read-back excluded. RSS recorded before validation, using Linux process VmHWM to exclude inherited pre-exec parent peaks (getrusage fallback elsewhere). A shared source fixture is generated once in the parent before all workers; CrabXL imports the validation engine only after timing and RSS collection. Temp spools sampled every 5 ms; output ZIP excluded. All values, coordinate order, format codes and cleanup verified. No native competitor or published-version performance claim.",
        "samples": samples,
        "summary": summary,
    }
    args.output.write_text(json.dumps(record, indent=2) + "\n")
    directory.cleanup()


if __name__ == "__main__":
    main()
