"""Same-call owned data-table creation/property/save comparison."""

import argparse
import importlib
import json
import platform
import statistics
import sys
import tempfile
import tomllib
from pathlib import Path

from array_calls import ROOT, measure

LITERALS = ["", "false", "invalid", "0", " true ", "quoted & value"]
FIELDS = ["ca", "dt2D", "dtr", "del1", "del2"]


def worker(engine, count, output):
    module = importlib.import_module(engine)
    DataTableFormula = importlib.import_module(
        engine + ".worksheet.formula"
    ).DataTableFormula
    book = module.Workbook()
    sheet = book.active
    for row in range(1, count + 1):
        literal = LITERALS[(row - 1) % len(LITERALS)]
        sheet.cell(
            row, 1, DataTableFormula(f"A{row}:B{row}", **dict.fromkeys(FIELDS, literal))
        )
        value = sheet.cell(row, 1).value
        assert all(getattr(value, name) == literal for name in FIELDS)
    book.save(output)
    book.close()
    print(count)


def verify(path, count):
    import openpyxl

    book = openpyxl.load_workbook(path)
    for row in range(1, count + 1):
        value = book.active.cell(row, 1).value
        assert value.ref == f"A{row}:B{row}"
        literal = LITERALS[(row - 1) % len(LITERALS)]
        assert all(
            getattr(value, name) == (literal if literal else False) for name in FIELDS
        )
    book.close()
    cached = openpyxl.load_workbook(path, data_only=True)
    assert all(cached.active.cell(row, 1).value is None for row in range(1, count + 1))
    cached.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--measure", type=Path, required=True)
    args = parser.parse_args()
    import openpyxl

    assert openpyxl.__version__ == "3.1.5"
    report = {
        "reference": openpyxl.__version__,
        "core": tomllib.loads((ROOT / "Cargo.toml").read_text())["dependencies"][
            "crabxl"
        ]["rev"],
        "python": sys.version,
        "platform": platform.platform(),
        "measurement": "One warmup plus five rotating serial cold-process wall/CPU/RSS samples including Python/import baseline; build/public readback excluded; temp sampling 25ms with cleanup checked",
        "semantics": "Identical owned Workbook/cell/DataTableFormula/property/save calls verify five raw flags before save and ordinary/cache readback outside timing. Six repeating literals include empty, Boolean spellings, opaque/whitespace/escaped strings. No prior adapter accepts every opaque flag and loaded-property edit, so no equivalent prior API benchmark is fabricated. No native competitor overlap or full workbook parity is claimed.",
        "cases": [],
    }
    with tempfile.TemporaryDirectory() as name:
        root = Path(name)
        temporary = root / "temporary"
        temporary.mkdir()
        for count in (1000, 10000):
            paths = {
                engine: root / f"{engine}.xlsx" for engine in ("crabxl", "openpyxl")
            }
            commands = {
                engine: [
                    args.measure.resolve(),
                    sys.executable,
                    __file__,
                    "--worker",
                    engine,
                    count,
                    paths[engine],
                ]
                for engine in paths
            }
            samples = {engine: [] for engine in paths}
            for engine, command in commands.items():
                output, _ = measure(command, temporary)
                assert output == str(count)
                verify(paths[engine], count)
            for iteration in range(5):
                for engine in (
                    ("crabxl", "openpyxl")
                    if iteration % 2 == 0
                    else ("openpyxl", "crabxl")
                ):
                    output, sample = measure(commands[engine], temporary)
                    assert output == str(count)
                    verify(paths[engine], count)
                    sample["output_bytes"] = paths[engine].stat().st_size
                    samples[engine].append(sample)
            report["cases"].append(
                {
                    "cells": count,
                    "samples": samples,
                    "medians": {
                        engine: {
                            key: statistics.median(sample[key] for sample in values)
                            for key in (
                                "seconds",
                                "cpu_seconds",
                                "peak_rss_kib",
                                "sampled_temp_peak_bytes",
                            )
                        }
                        for engine, values in samples.items()
                    },
                }
            )
            (ROOT / "benchmarks/results/table-calls.json").write_text(
                json.dumps(report, indent=2) + "\n"
            )
            print(count, report["cases"][-1]["medians"], flush=True)


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--worker":
        worker(sys.argv[2], int(sys.argv[3]), sys.argv[4])
    else:
        main()
