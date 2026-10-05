"""Measure equivalent scalar optimized modes in isolated serial processes."""

import argparse
import hashlib
import json
import platform
import statistics
import subprocess
import tempfile
from pathlib import Path


def worker(engine, operation, rows, path):
    module = __import__(engine if engine != "python-calamine" else "python_calamine")
    spool_peak = None
    options = {"max_memory_bytes": 1024**3} if engine == "crabxl" else {}
    if operation == "read":
        if engine == "python-calamine":
            book = module.CalamineWorkbook.from_path(str(path))
            values = book.get_sheet_by_name("Sheet").iter_rows()
        else:
            book = module.load_workbook(path, read_only=True, **options)
            values = book.active.iter_rows(values_only=True)
        count = total = 0
        for row in values:
            count += len(row)
            total += sum(row)
        book.close()
    else:
        if engine == "xlsxwriter":
            book = module.Workbook(str(path), {"constant_memory": True})
            sheet = book.add_worksheet("Sheet")
        else:
            book = module.Workbook(write_only=True, **options)
            sheet = book.create_sheet("Sheet")
        for row in range(rows):
            numbers = [row * 10 + column for column in range(10)]
            if engine == "xlsxwriter":
                sheet.write_row(row, 0, numbers)
            else:
                sheet.append(numbers)
        if engine == "crabxl":
            sheet.close()
            spool_peak = book._stream_writer.stats()[2]
        if engine != "xlsxwriter":
            book.save(path)
        book.close()
        count = rows * 10
        total = count * (count - 1) // 2
    print(
        json.dumps(
            {
                "cells": count,
                "checksum": total,
                "spool_peak_bytes": spool_peak,
                "output_bytes": path.stat().st_size if operation == "write" else None,
            }
        )
    )


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--worker", choices=["crabxl", "openpyxl", "python-calamine", "xlsxwriter"]
    )
    parser.add_argument("--operation", choices=["read", "write"])
    parser.add_argument("--rows", type=int, default=100000)
    parser.add_argument("--path", type=Path)
    parser.add_argument("--python", type=Path)
    parser.add_argument("--measure", type=Path)
    parser.add_argument("--runs", type=int, default=5)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    if args.worker:
        worker(args.worker, args.operation, args.rows, args.path)
        return
    import openpyxl

    report = {
        "semantics": "Scalar numeric read_only values and write_only append/save. "
        "One warmup, rotating serial cold processes including imports/cleanup. "
        "Generation and readback excluded; exact count/checksum verified. "
        "No full style/feature parity claim. Python 3.12, Linux x86_64.",
        "platform": platform.platform(),
        "packages": json.loads(
            subprocess.check_output(
                [
                    str(args.python),
                    "-c",
                    "import importlib.metadata,json,hashlib, pathlib,crabxl,crabxl._native,crabxl.optimized; "
                    "print(json.dumps(dict({n:importlib.metadata.version(n) for n in "
                    "['crabxl','openpyxl','python-calamine','xlsxwriter']}, "
                    "native_sha256=hashlib.sha256(pathlib.Path(crabxl._native.__file__).read_bytes()).hexdigest(), "
                    "python_sha256=hashlib.sha256(pathlib.Path(crabxl.__file__).read_bytes()).hexdigest(), "
                    "optimized_sha256=hashlib.sha256(pathlib.Path(crabxl.optimized.__file__).read_bytes()).hexdigest())))",
                ],
                text=True,
            )
        ),
        "core_revision": "7e9eb8db613d1f9fd921f8698c50f858c76009bb",
        "model_allowance_bytes": 1024**3,
        "cases": [],
    }
    with tempfile.TemporaryDirectory(prefix="crabxl-modes-benchmark-") as directory:
        for rows in [1000, 10000, args.rows]:
            source = Path(directory) / "source.xlsx"
            book = openpyxl.Workbook(write_only=True)
            sheet = book.create_sheet("Sheet")
            for row in range(rows):
                sheet.append([row * 10 + column for column in range(10)])
            book.save(source)
            for operation in ["read", "write"]:
                samples = {
                    name: []
                    for name in (
                        ["crabxl", "openpyxl", "python-calamine"]
                        if operation == "read"
                        else ["crabxl", "openpyxl", "xlsxwriter"]
                    )
                }
                for trial in range(args.runs + 1):
                    names = list(samples)
                    names = names[trial % len(names) :] + names[: trial % len(names)]
                    for engine in names:
                        path = (
                            source
                            if operation == "read"
                            else Path(directory) / f"{engine}.xlsx"
                        )
                        result = subprocess.run(
                            [
                                str(args.measure),
                                str(args.python),
                                str(Path(__file__).resolve()),
                                "--worker",
                                engine,
                                "--operation",
                                operation,
                                "--rows",
                                str(rows),
                                "--path",
                                str(path),
                            ],
                            check=True,
                            text=True,
                            capture_output=True,
                        )
                        output = json.loads(result.stdout)
                        cells = rows * 10
                        assert {key: output[key] for key in ["cells", "checksum"]} == {
                            "cells": cells,
                            "checksum": cells * (cells - 1) // 2,
                        }
                        sample = json.loads(result.stderr.split("MEASURE ")[-1])
                        sample.update(output)
                        if operation == "write":
                            check = openpyxl.load_workbook(path, read_only=True)
                            actual_count = actual_total = 0
                            for row in check.active.values:
                                actual_count += len(row)
                                actual_total += sum(row)
                            check.close()
                            assert (actual_count, actual_total) == (
                                cells,
                                cells * (cells - 1) // 2,
                            )
                        if trial:
                            samples[engine].append(sample)
                        print(rows, operation, engine, sample["seconds"], flush=True)
                report["cases"].append(
                    {
                        "rows": rows,
                        "columns": 10,
                        "input_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                        "operation": operation,
                        "samples": samples,
                        "median": {
                            name: {
                                key: statistics.median(item[key] for item in values)
                                for key in ["seconds", "peak_rss_kib"]
                            }
                            for name, values in samples.items()
                        },
                    }
                )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, indent=2) + "\n")


if __name__ == "__main__":
    main()
