"""Same-call temporal replacement and repeated owned save comparison."""

import argparse
import importlib
import json
import platform
import statistics
import sys
import tempfile
import tomllib
from datetime import date, datetime, time, timedelta
from pathlib import Path

from array_calls import ROOT, measure

INITIAL = [
    date(2024, 1, 2),
    datetime(2024, 1, 2, 3, 4, 5),
    time(3, 4, 5),
    timedelta(days=2, seconds=3),
]
REPLACEMENTS = [
    date(2024, 2, 3),
    datetime(2024, 2, 3, 4, 5, 6),
    time(4, 5, 6),
    timedelta(days=3, seconds=4),
    42,
]


def worker(engine, rows, output):
    module = importlib.import_module(engine)
    book = module.Workbook()
    sheet = book.active
    for row in range(1, rows + 1):
        for column, replacement in enumerate(REPLACEMENTS, 1):
            cell = sheet.cell(row, column)
            cell.value = INITIAL[(row - 1) % len(INITIAL)]
            cell.value = replacement
            assert cell.value == replacement
    book.save(output)
    book.save(output.with_name("repeat-" + output.name))
    book.close()
    print(rows * 5)


def record(cell):
    return type(cell.value).__name__, str(cell.value), cell.number_format


def records(path):
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True)
    values = [record(cell) for row in book.active for cell in row]
    book.close()
    return values


def verify(path, rows, expected):
    import openpyxl

    for saved in (path, path.with_name("repeat-" + path.name)):
        book = openpyxl.load_workbook(saved, read_only=True)
        count = 0
        for row in book.active:
            for cell in row:
                assert record(cell) == expected[count % len(expected)]
                count += 1
        assert count == rows * 5
        book.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--measure", type=Path, required=True)
    args = parser.parse_args()
    import openpyxl

    assert openpyxl.__version__ == "3.1.5"
    report = {
        "reference": openpyxl.__version__,
        "core": tomllib.loads((ROOT / "Cargo.toml").read_text())["dependencies"][
            "crabxl"
        ]["rev"],
        "python": sys.version,
        "platform": platform.platform(),
        "measurement": "One warmup plus five rotating serial cold-process wall/CPU/RSS samples including Python/import baseline; build/public readback excluded; temp sampling 25ms with cleanup checked",
        "semantics": "Identical owned Workbook/cell/value/getter calls assign four initial temporal kinds, replace with four kinds or a numeric value, then save twice. Existing formats survive replacement. Both saved outputs verify every value/type/format outside timing. No prior adapter retained these temporal assignment formats, so no equivalent prior feature baseline is fabricated. No native competitor or full style API claim.",
        "cases": [],
    }
    with tempfile.TemporaryDirectory() as name:
        root = Path(name)
        temporary = root / "temporary"
        temporary.mkdir()
        reference = root / "reference.xlsx"
        worker("openpyxl", 4, reference)
        expected = records(reference)
        for rows in (1000, 10000):
            paths = {
                engine: root / f"{engine}.xlsx" for engine in ("crabxl", "openpyxl")
            }
            commands = {
                engine: [
                    args.measure.resolve(),
                    sys.executable,
                    __file__,
                    "--worker",
                    engine,
                    rows,
                    paths[engine],
                ]
                for engine in paths
            }
            samples = {engine: [] for engine in paths}
            for engine, command in commands.items():
                output, _ = measure(command, temporary)
                assert output == str(rows * 5)
                verify(paths[engine], rows, expected)
            for iteration in range(5):
                for engine in (
                    ("crabxl", "openpyxl")
                    if iteration % 2 == 0
                    else ("openpyxl", "crabxl")
                ):
                    output, sample = measure(commands[engine], temporary)
                    assert output == str(rows * 5)
                    verify(paths[engine], rows, expected)
                    sample["output_bytes"] = (
                        paths[engine].stat().st_size
                        + paths[engine]
                        .with_name("repeat-" + paths[engine].name)
                        .stat()
                        .st_size
                    )
                    samples[engine].append(sample)
            report["cases"].append(
                {
                    "cells": rows * 5,
                    "saves": 2,
                    "samples": samples,
                    "medians": {
                        engine: {
                            key: statistics.median(sample[key] for sample in values)
                            for key in (
                                "seconds",
                                "cpu_seconds",
                                "peak_rss_kib",
                                "sampled_temp_peak_bytes",
                            )
                        }
                        for engine, values in samples.items()
                    },
                }
            )
            (ROOT / "benchmarks/results/temporal-calls.json").write_text(
                json.dumps(report, indent=2) + "\n"
            )
            print(rows, report["cases"][-1]["medians"], flush=True)


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--worker":
        worker(sys.argv[2], int(sys.argv[3]), Path(sys.argv[4]))
    else:
        main()
