"""Same-call array property/create/save comparison against the public reference."""

import argparse
import importlib
import json
import os
import platform
import statistics
import subprocess
import sys
import tempfile
import time
from pathlib import Path

TEXTS = [None, "", "=1", "abc", "==1", "\u03b1x"]
ROOT = Path(__file__).resolve().parents[1]


def worker(engine, count, output):
    module = importlib.import_module(engine)
    ArrayFormula = importlib.import_module(engine + ".worksheet.formula").ArrayFormula
    book = module.Workbook()
    sheet = book.active
    for row in range(1, count + 1):
        text = TEXTS[(row - 1) % len(TEXTS)]
        sheet.cell(row, 1, ArrayFormula(f"A{row}", text))
        assert sheet.cell(row, 1).value.text == text
    book.save(output)
    book.close()
    print(count)


def verify(path, count):
    import openpyxl

    book = openpyxl.load_workbook(path)
    for row in range(1, count + 1):
        value = book.active.cell(row, 1).value
        assert value.ref == f"A{row}"
        assert value.text == "=" + (TEXTS[(row - 1) % len(TEXTS)] or "")[1:]
    book.close()
    cached = openpyxl.load_workbook(path, data_only=True)
    assert all(cached.active.cell(row, 1).value is None for row in range(1, count + 1))
    cached.close()


def measure(command, directory):
    process = subprocess.Popen(
        list(map(str, command)),
        env=dict(os.environ, TMPDIR=str(directory)),
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    peak = 0
    while process.poll() is None:
        current = 0
        for path in directory.iterdir():
            try:
                current += path.stat().st_size
            except FileNotFoundError:
                pass
        peak = max(peak, current)
        time.sleep(0.025)
    output, errors = process.communicate()
    if process.returncode:
        raise RuntimeError(errors)
    sample = json.loads(errors.split("MEASURE ")[-1])
    sample["sampled_temp_peak_bytes"] = peak
    assert not list(directory.iterdir()), "Temporary-file leak"
    return output.strip(), sample


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--measure",
        type=Path,
        required=True,
        help="Compiled crabxl benchmarks/measure.c helper",
    )
    parser.add_argument(
        "--output", type=Path, default=ROOT / "benchmarks/results/array-calls.json"
    )
    args = parser.parse_args()
    import openpyxl

    assert openpyxl.__version__ == "3.1.5"
    report = {
        "reference": openpyxl.__version__,
        "core": __import__("tomllib").loads((ROOT / "Cargo.toml").read_text())[
            "dependencies"
        ]["crabxl"]["rev"],
        "platform": platform.platform(),
        "python": sys.version,
        "measurement": "One warmup and five rotating serial cold-process wall/CPU/RSS samples including Python/import baseline; build and public readback excluded; temporary files sampled every 25ms with cleanup checked",
        "semantics": "Identical Workbook/cell/ArrayFormula/property/save calls retain an owned sheet in both engines. None, empty, equals and non-equals/Unicode prefixes are validated before save; ordinary/data-only readback is verified outside timing. No reference runtime fallback in adapter; this is supported array-call overlap, not full workbook parity. No prior adapter pin had all these property semantics, so no equivalent old baseline is fabricated.",
        "cases": [],
    }
    with tempfile.TemporaryDirectory() as name:
        root = Path(name)
        temporary = root / "temporary"
        temporary.mkdir()
        for count in (1000, 10000):
            paths = {
                engine: root / f"{engine}.xlsx" for engine in ("crabxl", "openpyxl")
            }
            commands = {
                engine: [
                    args.measure.resolve(),
                    sys.executable,
                    __file__,
                    "--worker",
                    engine,
                    count,
                    paths[engine],
                ]
                for engine in paths
            }
            samples = {engine: [] for engine in paths}
            for engine, command in commands.items():
                output, _ = measure(command, temporary)
                assert output == str(count)
                verify(paths[engine], count)
            for iteration in range(5):
                for engine in (
                    ("crabxl", "openpyxl")
                    if iteration % 2 == 0
                    else ("openpyxl", "crabxl")
                ):
                    output, sample = measure(commands[engine], temporary)
                    assert output == str(count)
                    verify(paths[engine], count)
                    sample["output_bytes"] = paths[engine].stat().st_size
                    samples[engine].append(sample)
            report["cases"].append(
                {
                    "cells": count,
                    "samples": samples,
                    "medians": {
                        engine: {
                            key: statistics.median(sample[key] for sample in values)
                            for key in (
                                "seconds",
                                "cpu_seconds",
                                "peak_rss_kib",
                                "sampled_temp_peak_bytes",
                            )
                        }
                        for engine, values in samples.items()
                    },
                }
            )
            args.output.write_text(json.dumps(report, indent=2) + "\n")
            print(count, report["cases"][-1]["medians"], flush=True)


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--worker":
        worker(sys.argv[2], int(sys.argv[3]), sys.argv[4])
    else:
        main()
