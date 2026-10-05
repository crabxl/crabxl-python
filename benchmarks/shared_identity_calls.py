"""Same-call shared-formula loading with literal group identity."""
import argparse
import hashlib
import importlib
import json
from pathlib import Path
import statistics
import sys
import tempfile
import zipfile
from array_calls import measure

ROOT = Path(__file__).resolve().parents[1]


def fixture(path, count):
    import openpyxl
    book = openpyxl.Workbook()
    book.save(path)
    book.close()
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    rows = []
    for row in range(1, count + 1):
        expression = 'A1+1' if row == 1 else ''
        rows.append(f'<row r="{row}"><c r="A{row}"><f t="shared" si="group &amp; value">{expression}</f><v>{row+1}</v></c></row>')
    parts['xl/worksheets/sheet1.xml'] = ('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>' + ''.join(rows) + '</sheetData></worksheet>').encode()
    with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def worker(engine, path, count, cached):
    module = importlib.import_module(engine)
    book = module.load_workbook(path, data_only=cached)
    digest = hashlib.sha256()
    for row, values in enumerate(book.active.iter_rows(min_row=1, max_row=count, max_col=1, values_only=True), 1):
        value = values[0]
        assert value == (row + 1 if cached else f'=A{row}+1')
        digest.update((str(value) + '\n').encode())
    book.close()
    print(digest.hexdigest())


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--measure', type=Path, required=True)
    parser.add_argument('--output', type=Path, default=ROOT / 'benchmarks/results/shared-identity-calls.json')
    args = parser.parse_args()
    import openpyxl
    assert openpyxl.__version__ == '3.1.5'
    report = {
        'core': __import__('tomllib').loads((ROOT / 'Cargo.toml').read_text())['dependencies']['crabxl']['rev'],
        'reference': openpyxl.__version__, 'python': sys.version,
        'measurement': 'One warmup and five alternating serial cold-process samples, including imports, load, iteration, identical per-value assertions and SHA256; fixture creation excluded. Kernel peak RSS and wall/CPU include runtime. Temporary bytes sampled every 25ms and cleanup checked; short-lived files may be missed.',
        'semantics': 'Identical ordinary load_workbook(data_only=...) and iter_rows(values_only=True) calls; loaded/materialized adapter ownership, not read-only streaming. Literal shared ID group & value matches exactly. Every formula/cache is checked. No prior adapter pin supported this full literal-ID workload; no prior equivalent baseline is claimed. Rust native/calculation speed is not inferred from adapter measurements.',
        'cases': [],
    }
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        temporary = root / 'temporary'
        temporary.mkdir()
        for count in (5000, 50000):
            path = root / 'source.xlsx'
            fixture(path, count)
            for cached in (False, True):
                commands = {engine: [args.measure.resolve(), sys.executable, __file__, '--worker', engine, path, count, int(cached)] for engine in ('crabxl', 'openpyxl')}
                samples = {engine: [] for engine in commands}
                hashes = {engine: measure(command, temporary)[0] for engine, command in commands.items()}
                assert hashes['crabxl'] == hashes['openpyxl']
                for iteration in range(5):
                    for engine in (('crabxl', 'openpyxl') if iteration % 2 == 0 else ('openpyxl', 'crabxl')):
                        digest, sample = measure(commands[engine], temporary)
                        assert digest == hashes[engine]
                        samples[engine].append(sample)
                report['cases'].append({'cells': count, 'data_only': cached, 'sha256': hashes['crabxl'], 'source_bytes': path.stat().st_size, 'samples': samples, 'medians': {engine: {key: statistics.median(item[key] for item in values) for key in ('seconds', 'cpu_seconds', 'peak_rss_kib', 'sampled_temp_peak_bytes')} for engine, values in samples.items()}})
                args.output.write_text(json.dumps(report, indent=2) + '\n')


if __name__ == '__main__':
    if len(sys.argv) > 1 and sys.argv[1] == '--worker':
        worker(sys.argv[2], sys.argv[3], int(sys.argv[4]), bool(int(sys.argv[5])))
    else:
        main()
