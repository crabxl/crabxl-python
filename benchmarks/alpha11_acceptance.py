"""Fresh-process A11 feature benchmark with independent saved-output checks.

Usage: python alpha11_acceptance.py ENGINE MODE COUNT DIRECTORY [SOURCE]
Modes: generate, owned, loaded, write_only. Builds/tests must not overlap samples.
"""

import ctypes
import importlib
import json
import os
import sys
import tempfile
import threading
import time
from pathlib import Path


def peak_rss_bytes():
    if os.name != "nt":
        import resource

        value = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
        return value if sys.platform == "darwin" else value * 1024

    class Counters(ctypes.Structure):
        _fields_ = [("cb", ctypes.c_ulong), ("faults", ctypes.c_ulong)] + [
            (name, ctypes.c_size_t)
            for name in (
                "peak",
                "working",
                "pool_peak",
                "pool",
                "nonpaged_peak",
                "nonpaged",
                "pagefile",
                "peak_pagefile",
            )
        ]

    counters = Counters()
    counters.cb = ctypes.sizeof(counters)
    if not ctypes.windll.psapi.GetProcessMemoryInfo(
        ctypes.c_void_p(-1), ctypes.byref(counters), counters.cb
    ):
        raise ctypes.WinError()
    return counters.peak


def populate(engine, book, sheet, count, streaming):
    rich = importlib.import_module(engine.__name__ + ".cell.rich_text")
    font = importlib.import_module(engine.__name__ + ".cell.text").InlineFont
    book.add_named_style(
        engine.styles.NamedStyle(name="Heading", font=engine.styles.Font(b=True))
    )
    sheet.column_dimensions["A"].width = 18
    for index in range(count):
        value = rich.CellRichText("row ", rich.TextBlock(font(b=True), str(index)))
        if streaming:
            numeric = engine.cell.WriteOnlyCell(sheet, index)
            numeric.style = "Heading"
            linked = engine.cell.WriteOnlyCell(sheet, "label")
            linked.hyperlink = f"https://example.org/{index}"
            sheet.append([numeric, value, None, linked])
        else:
            row = index * 3 + 1
            sheet.cell(row, 1, index).style = "Heading"
            sheet.cell(row, 2, value)
            sheet.cell(row, 4, "label").hyperlink = f"https://example.org/{index}"
            sheet.cell(row, 6, "merge")
            sheet.merge_cells(
                start_row=row, start_column=6, end_row=row + 1, end_column=7
            )


def main():
    engine_name, mode, raw_count, root = sys.argv[1:5]
    count = int(raw_count)
    root = Path(root).resolve()
    root.mkdir(parents=True, exist_ok=True)
    if mode == "generate":
        import openpyxl

        book = openpyxl.Workbook()
        populate(openpyxl, book, book.active, count, False)
        book.save(root / f"source-{count}.xlsx")
        book.close()
        return
    with tempfile.TemporaryDirectory(prefix="a11-", dir=root) as folder:
        directory = Path(folder)
        os.environ["TEMP"] = os.environ["TMP"] = str(directory)
        tempfile.tempdir = str(directory)
        output = directory / "output.xlsx"
        stop = threading.Event()
        peak_temp = [0]

        def sample():
            while not stop.is_set():
                total = 0
                for path in directory.rglob("*"):
                    if path == output:
                        continue
                    try:
                        if path.is_file():
                            total += path.stat().st_size
                    except FileNotFoundError:
                        pass
                peak_temp[0] = max(peak_temp[0], total)
                stop.wait(0.01)

        monitor = threading.Thread(target=sample, daemon=True)
        monitor.start()
        cpu = time.process_time()
        started = time.perf_counter()
        engine = importlib.import_module(engine_name)
        if mode == "loaded":
            book = engine.load_workbook(sys.argv[5], rich_text=True)
            sheet = book.active
        else:
            book = engine.Workbook(write_only=mode == "write_only")
            sheet = book.create_sheet() if mode == "write_only" else book.active
            populate(engine, book, sheet, count, mode == "write_only")
        if mode != "write_only":
            sheet.insert_rows(1)
            sheet.move_range(f"A1:D{count * 3}", rows=1)
        prepared = time.perf_counter()
        book.save(output)
        if mode != "write_only":
            book.save(output)
        book.close()
        elapsed = time.perf_counter() - started
        cpu = time.process_time() - cpu
        peak = peak_rss_bytes()
        stop.set()
        monitor.join()
        leftover = [path.name for path in directory.iterdir() if path != output]
        assert not leftover, leftover
        import openpyxl

        verified = openpyxl.load_workbook(output, rich_text=True)
        target = verified.active
        assert (
            sum(cell[0].value or 0 for cell in target.iter_rows(max_col=1))
            == count * (count - 1) // 2
        )
        last = count if mode == "write_only" else (count - 1) * 3 + 3
        assert target.cell(last, 1).value == count - 1
        assert str(target.cell(last, 2).value) == f"row {count - 1}"
        assert target.cell(last, 1).font.b
        link_row = count if mode == "write_only" else (count - 1) * 3 + 1
        assert (
            target.cell(link_row, 4).hyperlink.target
            == f"https://example.org/{count - 1}"
        )
        if mode != "write_only":
            assert len(target.merged_cells.ranges) == count
        verified.close()
        print(
            json.dumps(
                dict(
                    engine=engine_name,
                    mode=mode,
                    count=count,
                    elapsed_seconds=elapsed,
                    prepare_seconds=prepared - started,
                    cpu_seconds=cpu,
                    peak_rss_bytes=peak,
                    peak_temp_sampled_bytes=peak_temp[0],
                    temp_cleanup=True,
                    output_bytes=output.stat().st_size,
                    checksum=count * (count - 1) // 2,
                )
            )
        )


if __name__ == "__main__":
    main()
